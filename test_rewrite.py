"""
test_rewrite.py
===============
Offline tests voor de deterministische lagen: de code-check (rewrite_checks.py), de
besluiten-parser (besluiten.py) en de CMS-output (rewrite_output.py).
Geen API-key nodig. Draai met `python test_rewrite.py` of `pytest test_rewrite.py`.

De LLM-classificatie in besluiten.py wordt hier bewust NIET getest — die vraagt een
API-call. Wat hier wél gegarandeerd moet kloppen is de laag eronder: de structurele
splitsing van `actie_besluit` en de uitlijning met `actualiteit_actie`.
"""

from __future__ import annotations

import contextlib
import copy
import io
import json
import os
import re
import tempfile
import time
from types import SimpleNamespace

import besluiten as bes
import drive_upload as drive
import rewrite_checks as checks
import rewrite_output as uit
import rewrite_trainings as rw
import sjabloon
from rewrite_checks import HARD, FLAG, check_rewrite, hard_fails, flags


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def vul(n: int, *woorden: str) -> str:
    """`n` vulwoorden om een lengtecheck op een exact woordaantal te zetten.

    Wisselt tussen de opgegeven woorden af, want twee keer hetzelfde woord achter elkaar is
    sinds `check_generic` een harde fout ("ervaar je hoe hoe ..."). Een fixture die 55 keer
    "data" achter elkaar zet, test dan niet meer wat hij bedoelt te testen.
    """
    keuze = woorden or ("data", "inzicht")
    return " ".join(keuze[i % len(keuze)] for i in range(n))


def _good_rewrite() -> dict:
    """Een concept dat ALLE harde checks haalt (0 hard-fails)."""
    return {
        # "kunnen" staat er bewust in: dit is het concept dat alle regels haalt, dus het hoort
        # ook het lerende aspect te demonstreren (schrijfspec Sectie 0.15).
        "overzicht": "Wil je data kunnen " + vul(55) + "?",   # 59 woorden
        "inleiding": vul(195, "onderwerp", "thema"),             # 195 woorden
        "modules": {"modules": [
            {"titel": "Module een", "bullets": ["Onderdeel a", "Onderdeel b", "Onderdeel c"]},
            {"titel": "Module twee", "bullets": ["Onderdeel a", "Onderdeel b", "Onderdeel c", "Onderdeel d"]},
            {"titel": "Module drie", "bullets": ["Onderdeel a", "Onderdeel b", "Onderdeel c"]},
            {"titel": "Module vier", "bullets": ["Onderdeel a", "Onderdeel b", "Onderdeel c", "Onderdeel d", "Onderdeel e"]},
        ]},
        "aanpak_invulling": "je datagedreven keuzes maakt",
        "doelgroep": "Deze training is bedoeld voor iedereen die met data betere keuzes wil maken.",
        "voorkennis": "Specifieke voorkennis voor het volgen van deze training is niet noodzakelijk.",
        "doelen": ["Heldere dashboards te bouwen voor je team",
                   "Ruwe data op te schonen en samen te voegen",
                   "Terugkerende trends te analyseren",
                   "Resultaten te presenteren aan het team"],
        "vervolgstappen_titels": ["Training Power BI"],
        "kortste_omschrijving": ("Wil je slimmer met data kunnen werken? Na deze training weet "
                                 "je hoe je betere keuzes onderbouwt."),
        "nieuwe_titel": "Training Data-analyse",
    }


_CTX = {"catalog_titles": {"Training Power BI", "Training T-SQL"}, "naam": "Training Data-analyse"}


def _codes(issues, severity=None):
    return {i.code for i in issues if severity is None or i.severity == severity}


def _codes_in(issues, section, severity=None):
    """Codes van één kopje. Nodig waar de fixture elders al hetzelfde signaal geeft --
    `_good_rewrite` bestaat uit herhaalde woorden zonder interpunctie, dus de Inleiding is
    formeel één zin van 195 woorden."""
    return _codes([i for i in issues if i.section == section], severity)


# ---------------------------------------------------------------------------
# Baseline: het goede concept haalt alle harde checks
# ---------------------------------------------------------------------------

def test_good_rewrite_has_no_hard_fails():
    issues = check_rewrite(_good_rewrite(), _CTX)
    hf = hard_fails(issues)
    assert hf == [], f"onverwachte hard-fails: {[str(i) for i in hf]}"


# ---------------------------------------------------------------------------
# Harde checks per kopje
# ---------------------------------------------------------------------------

def test_korte_te_kort():
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je " + vul(10) + "?"
    assert "lengte_woorden" in _codes(check_rewrite(rw, _CTX), HARD)


def test_korte_verkeerde_opening():
    rw = _good_rewrite()
    rw["overzicht"] = "Deze training " + vul(57) + "."
    assert "opening" in _codes(check_rewrite(rw, _CTX), HARD)


def test_algemene_te_lang():
    rw = _good_rewrite()
    rw["inleiding"] = vul(280, "onderwerp", "thema")
    assert "lengte_woorden" in _codes(check_rewrite(rw, _CTX), HARD)


# ---------------------------------------------------------------------------
# Lengte is een richtlijn met een vangrail: net eroverheen mag, ver eroverheen niet
# ---------------------------------------------------------------------------

def test_korte_net_buiten_richtlijn_is_flag():
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je " + vul(90) + "?"   # 92 woorden: buiten 55-80, binnen 45-110
    issues = check_rewrite(rw, _CTX)
    assert "lengte_woorden" not in _codes(issues, HARD)
    assert "lengte_richtlijn" in _codes(issues, FLAG)


def test_een_ruimer_overzicht_mag_gewoon():
    """De band ging in ronde 2 van 55-65 naar 55-80: op de eigen catalogus is de mediaan 64.

    "Lengtebeperking is geen doel op zich" -- een Overzicht van 75 woorden dat de kern compleet
    maakt hoort niet eens een flag op te leveren.
    """
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je " + vul(73) + "?"   # 75 woorden
    issues = check_rewrite(rw, _CTX)
    assert "lengte_richtlijn" not in _codes_in(issues, "overzicht", FLAG)
    assert "lengte_woorden" not in _codes(issues, HARD)


def test_korte_buiten_vangrail_is_hard():
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je " + vul(120) + "?"
    assert "lengte_woorden" in _codes(check_rewrite(rw, _CTX), HARD)


def test_algemene_net_buiten_richtlijn_is_flag():
    rw = _good_rewrite()
    rw["inleiding"] = vul(228, "onderwerp", "thema")
    issues = check_rewrite(rw, _CTX)
    assert "lengte_woorden" not in _codes(issues, HARD)
    assert "lengte_richtlijn" in _codes(issues, FLAG)


def test_inleidingsband_schuift_met_dagen():
    """Een training van vijf dagen mag een langere Inleiding hebben dan een van één dag."""
    rw = _good_rewrite()
    rw["inleiding"] = vul(228, "onderwerp", "thema")
    lang = check_rewrite(rw, dict(_CTX, dagen=5))
    kort = check_rewrite(rw, dict(_CTX, dagen=1))
    assert "lengte_richtlijn" not in _codes(lang, FLAG)   # binnen de band voor 5 dagen
    assert "lengte_richtlijn" in _codes(kort, FLAG)       # te lang voor een eendaagse
    # de vangrail schuift mee, maar niet oneindig
    rw["inleiding"] = vul(265, "onderwerp", "thema")
    assert "lengte_woorden" not in _codes(check_rewrite(rw, dict(_CTX, dagen=5)), HARD)
    assert "lengte_woorden" in _codes(check_rewrite(rw, dict(_CTX, dagen=2)), HARD)


def test_lange_zin_is_flag_geen_hardfail():
    rw = _good_rewrite()
    # 40 woorden in één zin: ver boven de richtlijn van ±20, dus een signaal -- maar de
    # schrijver hoort er niet voor terug te moeten.
    rw["doelgroep"] = "Deze training is bedoeld voor " + " ".join(["iedereen"] * 37) + "."
    issues = check_rewrite(rw, _CTX)
    assert "zin_lang" in _codes_in(issues, "doelgroep", FLAG)
    assert "zin_lang" not in _codes(issues, HARD)


def test_zin_boven_de_richtlijn_is_geen_flag():
    """25 woorden is langer dan ±20 en moet gewoon mogen: de richtlijn is geen plafond."""
    rw = _good_rewrite()
    rw["voorkennis"] = "Je werkt " + " ".join(["dagelijks"] * 23) + "."
    assert "zin_lang" not in _codes_in(check_rewrite(rw, _CTX), "voorkennis", FLAG)


def test_zinlengte_kijkt_niet_naar_bullets():
    """Bullets zijn geen zinnen; een lange module-bullet is geen zinlengte-signaal."""
    rw = _good_rewrite()
    rw["modules"]["modules"][0]["bullets"][0] = " ".join(["onderdeel"] * 45)
    assert "zin_lang" not in _codes_in(check_rewrite(rw, _CTX), "modules", FLAG)


def test_overzichtsband_negeert_dagen():
    """Het Overzicht is de aanhaakalinea; die blijft even lang, hoe lang de training ook duurt."""
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je " + vul(90) + "?"
    for dagen in (1, 5):
        assert "lengte_richtlijn" in _codes(check_rewrite(rw, dict(_CTX, dagen=dagen)), FLAG)


def test_programma_te_weinig_modules():
    rw = _good_rewrite()
    rw["modules"]["modules"] = rw["modules"]["modules"][:3]
    assert "modules_aantal" in _codes(check_rewrite(rw, _CTX), HARD)


def test_programma_bullets_buiten_bereik():
    rw = _good_rewrite()
    rw["modules"]["modules"][0]["bullets"] = ["een", "twee"]  # 2 < 3
    assert "bullets_aantal" in _codes(check_rewrite(rw, _CTX), HARD)


def test_programma_bullets_geen_variatie():
    rw = _good_rewrite()
    for m in rw["modules"]["modules"]:
        m["bullets"] = ["a", "b", "c"]  # overal 3
    assert "bullets_variatie" in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelgroep_professionals():
    rw = _good_rewrite()
    rw["doelgroep"] = "Deze training is bedoeld voor professionals die met data werken."
    assert "professionals" in _codes(check_rewrite(rw, _CTX), HARD)


# ---------------------------------------------------------------------------
# Verboden woorden (humanisering_nl.md Sectie D)
# ---------------------------------------------------------------------------

def test_professionals_is_hard_in_elk_kopje():
    """Het verbod gold alleen in doelgroep; het geldt nu overal."""
    rw = _good_rewrite()
    rw["inleiding"] = rw["inleiding"].replace("onderwerp", "professionals", 1)
    assert "professionals" in _codes(check_rewrite(rw, _CTX), HARD)


def test_professionals_in_de_titel_blijft_toegestaan():
    """"Training PHP Professional" bestaat echt; die training moet zichzelf kunnen noemen."""
    rw = _good_rewrite()
    rw["nieuwe_titel"] = "Training PHP Professional"
    ctx = dict(_CTX, naam="Training PHP Professional")
    assert "professionals" not in _codes(check_rewrite(rw, ctx), HARD)


def test_professionals_degradeert_naar_flag_bij_professional_titel():
    rw = _good_rewrite()
    rw["doelgroep"] = "Deze training is bedoeld voor iedereen die de PHP Professional-stof wil beheersen."
    ctx = dict(_CTX, naam="Training PHP Professional")
    codes = _codes(check_rewrite(rw, ctx))
    assert "professionals" in codes
    assert "professionals" not in _codes(check_rewrite(rw, ctx), HARD)


def test_bezig_met_is_hard():
    rw = _good_rewrite()
    rw["overzicht"] = rw["overzicht"].replace(
        "Wil je", "Wil je weten waarmee je houdt je bezig met data en", 1)
    assert "bezig_met" in _codes(check_rewrite(rw, _CTX), HARD)


def test_meeting_is_flag_geen_hard_fail():
    """Vakterm in Scrum/Agile/Teams-trainingen, dus signaleren en niet blokkeren."""
    rw = _good_rewrite()
    rw["doelgroep"] = "Deze training is bedoeld voor iedereen die elke meeting beter wil voorbereiden."
    assert "meeting" in _codes(check_rewrite(rw, _CTX), FLAG)
    assert "meeting" not in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelgroep_verkeerde_opening():
    rw = _good_rewrite()
    rw["doelgroep"] = "Voor iedereen die met data werkt."
    assert "opening" in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelen_verkeerd_aantal():
    rw = _good_rewrite()
    rw["doelen"] = rw["doelen"][:2]
    assert "aantal" in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelen_kleine_letter():
    rw = _good_rewrite()
    rw["doelen"][0] = "heldere dashboards te bouwen"   # wel te-infinitief, geen hoofdletter
    codes = _codes(check_rewrite(rw, _CTX), HARD)
    assert "hoofdletter" in codes
    assert "geen_te_infinitief" not in codes


def test_doelen_zonder_te_infinitief():
    """De kale infinitief ('Dashboards bouwen') loopt niet door op de vaste introzin."""
    rw = _good_rewrite()
    rw["doelen"][0] = "Dashboards bouwen die de juiste vraag beantwoorden"
    assert "geen_te_infinitief" in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelen_gesplitste_te_infinitief_is_goed():
    """'voor te bereiden' / 'uit te oefenen' zijn geldige te-infinitieven."""
    rw = _good_rewrite()
    rw["doelen"][0] = "Jezelf voor te bereiden op een gesprek met je team"
    rw["doelen"][1] = "Invloed uit te oefenen zonder formele macht"
    assert "geen_te_infinitief" not in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelen_onregelmatige_infinitief_is_goed():
    rw = _good_rewrite()
    rw["doelen"][0] = "Data te zien als basis voor besluitvorming"
    rw["doelen"][1] = "Om te gaan met tegenstrijdige belangen in het team"
    assert "geen_te_infinitief" not in _codes(check_rewrite(rw, _CTX), HARD)


def test_doelen_herhalen_in_staat_niet():
    """De introzin zegt al "ben je in staat om"; een bullet mag dat niet dubbelen."""
    rw = _good_rewrite()
    rw["doelen"][0] = "In staat te zijn om heldere dashboards te bouwen"
    assert "dubbel_in_staat" in _codes(check_rewrite(rw, _CTX), FLAG)


def test_doelen_inzicht_constructie_is_goed():
    """"Inzicht te krijgen in ..." is een aanbevolen causale constructie, geen fout."""
    rw = _good_rewrite()
    rw["doelen"][0] = "Inzicht te krijgen in de kosten van je datamodel"
    codes = _codes(check_rewrite(rw, _CTX))
    assert "geen_te_infinitief" not in codes
    assert "dubbel_in_staat" not in codes


def test_doelen_lopen_door_op_de_vaste_introzin():
    """Elke bullet moet als één zin achter sjabloon.DOELEN_INTRO te lezen zijn."""
    rw = _good_rewrite()
    assert sjabloon.DOELEN_INTRO.endswith("om:")
    for bullet in rw["doelen"]:
        zin = f"{sjabloon.DOELEN_INTRO[:-1]} {bullet[0].lower()}{bullet[1:]}"
        assert " te " in zin, zin


def test_kortste_te_lang():
    rw = _good_rewrite()
    rw["kortste_omschrijving"] = "Wil je " + "x" * 250 + "?"
    assert "lengte_tekens" in _codes(check_rewrite(rw, _CTX), HARD)


def test_kortste_verkeerde_opening():
    rw = _good_rewrite()
    rw["kortste_omschrijving"] = "Leer slimmer werken met data en tools."
    assert "opening" in _codes(check_rewrite(rw, _CTX), HARD)


# ---------------------------------------------------------------------------
# Generieke harde checks: placeholders, HTML, onbekende catalogus-titel
# ---------------------------------------------------------------------------

def test_placeholder_blijft_staan():
    rw = _good_rewrite()
    rw["aanpak_invulling"] = "je werkt met [....] in de praktijk"
    assert "placeholder" in _codes(check_rewrite(rw, _CTX), HARD)


def test_html_in_tekst():
    rw = _good_rewrite()
    rw["modules"]["modules"][0]["bullets"][0] = "<p>Onderdeel</p>"
    assert "html" in _codes(check_rewrite(rw, _CTX), HARD)


def test_onbekende_vervolgtraining_titel():
    rw = _good_rewrite()
    rw["vervolgstappen_titels"] = ["Training Bestaat Niet"]
    assert "titel_onbekend" in _codes(check_rewrite(rw, _CTX), HARD)


def test_ontbrekend_kopje():
    rw = _good_rewrite()
    rw["doelgroep"] = ""
    assert "ontbreekt" in _codes(check_rewrite(rw, _CTX), HARD)


# ---------------------------------------------------------------------------
# Soortwoord: niks heet nog een cursus of opleiding
# ---------------------------------------------------------------------------

def test_cursus_in_lopende_tekst_is_hardfail():
    rw = _good_rewrite()
    rw["inleiding"] = "Tijdens deze cursus " + vul(193, "onderwerp", "thema")
    assert "soortwoord" in _codes(check_rewrite(rw, _CTX), HARD)


def test_opleiding_in_module_bullet_is_hardfail():
    rw = _good_rewrite()
    rw["modules"]["modules"][0]["bullets"][0] = "Opzet van de opleiding"
    assert "soortwoord" in _codes(check_rewrite(rw, _CTX), HARD)


def test_cursus_in_titel_is_hardfail():
    rw = _good_rewrite()
    rw["nieuwe_titel"] = "Cursus XML"
    assert "soortwoord" in _codes(check_rewrite(rw, _CTX), HARD)


def test_examentraining_mag_wel():
    rw = _good_rewrite()
    rw["nieuwe_titel"] = "Examentraining DAMA-DMBOK CDMP"
    assert "soortwoord" not in _codes(check_rewrite(rw, _CTX), HARD)


# ---------------------------------------------------------------------------
# Titelnormalisatie (sjabloon.nieuwe_titel / vervang_soortwoord / vervolgtitel)
# ---------------------------------------------------------------------------

def test_nieuwe_titel_vervangt_verboden_soortwoord():
    assert sjabloon.nieuwe_titel("Opleiding PHP Professional") == "Training PHP Professional"
    assert sjabloon.nieuwe_titel("Cursus XML") == "Training XML"
    assert sjabloon.nieuwe_titel("Gebruikerscursus Sitecore") == "Training Sitecore"


def test_nieuwe_titel_laat_toegestane_soortwoorden_staan():
    for titel in ("Training Linux", "Masterclass PHP", "Workshop Storytelling",
                  "Examentraining CEH"):
        assert sjabloon.nieuwe_titel(titel) == titel


def test_nieuwe_titel_zet_voorvoegsel_bij_titel_zonder_soortwoord():
    assert sjabloon.nieuwe_titel("Excel") == "Training Excel"


def test_vervang_soortwoord_laat_niet_titels_met_rust():
    """In gekopieerde tekst staan regels die geen titel zijn; die niet aanraken."""
    zin = "Trainingen voor specifieke databasesystemen zoals PostgreSQL"
    assert sjabloon.vervang_soortwoord(zin) == zin
    assert sjabloon.vervang_soortwoord("Cursus PowerPoint") == "Training PowerPoint"


def test_vervolgtitel_haalt_het_voorvoegsel_weg():
    """In de Vervolgstappen-lijst is "Training" bij elke regel ruis."""
    assert sjabloon.vervolgtitel("Training Power BI") == "Power BI"
    assert sjabloon.vervolgtitel("Cursus PowerPoint") == "PowerPoint"
    assert sjabloon.vervolgtitel("Opleiding PHP Professional") == "PHP Professional"
    assert sjabloon.vervolgtitel("Power BI") == "Power BI"


def test_vervolgtitel_houdt_een_afwijkende_vorm():
    """Masterclass/workshop/examentraining zeggen iets over de vorm: die blijft staan."""
    for titel in ("Masterclass PHP", "Workshop Storytelling", "Examentraining CEH"):
        assert sjabloon.vervolgtitel(titel) == titel


def test_vervolgtitel_laat_regels_die_geen_titel_zijn_met_rust():
    zin = "Trainingen voor specifieke databasesystemen zoals PostgreSQL"
    assert sjabloon.vervolgtitel(zin) == zin
    # lopende zin met een verboden soortwoord: hooguit dát woord vervangen
    assert sjabloon.vervolgtitel("Cursus voor gevorderden") == "Training voor gevorderden"
    # een hoofdletter erna maakt het wél een titel
    assert sjabloon.vervolgtitel("Training Van Excel naar Power BI") == "Van Excel naar Power BI"


# ---------------------------------------------------------------------------
# Vervolgtrainingen: taxonomieboom + shortlist
# ---------------------------------------------------------------------------

def _mini_catalog() -> list[dict]:
    """Catalogus met net genoeg structuur om de shortlist te kunnen sturen.

    XSL en JavaScript delen geen woord maar wel een vakgebied; Active Directory deelt
    wél woorden met LDAP maar hangt in een ander domein. Precies de twee gevallen waarop
    keyword-overlap en boom afzonderlijk stukgaan.
    """
    rijen = [
        (1, "Cursus XSL", "Transformeer XML-documenten met stylesheets."),
        (2, "Cursus JavaScript", "Programmeer interactie in de browser."),
        (3, "Cursus Node.js", "Bouw serverapplicaties in de browsertaal."),
        (4, "Cursus LDAP", "Ontsluit centraal opgeslagen directory-gegevens via het netwerk."),
        (5, "Training Active Directory", "Beheer directory-gegevens en gebruikers centraal."),
        (6, "Training 5G Mobiele Communicatie", "Mobiele netwerken van de vijfde generatie."),
        (7, "Training Bloemschikken", "Niets met techniek te maken."),
        (8, "Opleiding C# Professional", "Professioneel programmeren in C-sharp."),
        (9, "Training C++ Professional", "Professioneel programmeren in C-plus-plus."),
        (10, "Training Claude CoWork", "Samenwerken met een AI-assistent."),
    ]
    return rw.catalog_uit_rijen([
        {"product_id": pid, "titel": t, "summary": s} for pid, t, s in rijen])


_MINI_BOOM = {
    "name": "Trainingscatalogus",
    "children": [
        {"name": "Software Development", "children": [
            {"name": "Web Development", "children": [
                {"name": "Cursus XSL"}, {"name": "Cursus JavaScript"}, {"name": "Cursus Node.js"},
            ]},
            {"name": "Programmeertalen", "children": [
                {"name": "Opleiding C# Professional"}, {"name": "Training C++ Professional"},
            ]},
        ]},
        {"name": "Cloud & Infrastructuur", "children": [
            {"name": "Netwerken & Connectiviteit", "children": [
                {"name": "Cursus LDAP"}, {"name": "Training 5G Mobiele Communicatie"},
            ]},
        ]},
        {"name": "Modern Workplace", "children": [
            {"name": "Identity & Access Management", "children": [
                {"name": "Training Active Directory"},
            ]},
        ]},
        # bestaat niet in de catalogus -> mag nooit in de index belanden
        {"name": "Data & Analytics", "children": [
            {"name": "Big Data", "children": [{"name": "Cursus Apache Pig"}]},
        ]},
    ],
}


def _mini_boom(catalog, boom=None):
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
        json.dump(boom if boom is not None else _MINI_BOOM, f)
        pad = f.name
    try:
        return rw.load_tree(catalog, pad)
    finally:
        os.unlink(pad)


def test_boom_bevat_alleen_bestaande_catalogustitels():
    """De veiligheidsgarantie: wat hier in staat, overleeft check_vervolgstappen.

    Elke titel buiten de catalogus is een HARD `titel_onbekend`, dus een blad zonder
    catalogusrij (vervallen aanbod als Apache Pig) mag niet als kandidaat naar buiten.
    """
    catalog = _mini_catalog()
    boom = _mini_boom(catalog)
    bekend = {t.strip().lower() for t in rw.catalog_titles(catalog)}
    assert boom["paden"], "de boom moet wél gevuld zijn"
    for sleutel in boom["paden"]:
        assert sleutel in bekend, sleutel
    assert "apache pig" not in boom["paden"]


def test_boom_koppelt_over_interpunctie_heen():
    """"Claude Co-Work" in de boom is dezelfde training als "Claude CoWork"."""
    catalog = _mini_catalog()
    boom = _mini_boom(catalog, {"name": "wortel", "children": [
        {"name": "Artificial Intelligence", "children": [
            {"name": "AI-assistenten", "children": [{"name": "Training Claude Co-Work"}]}]}]})
    assert "claude cowork" in boom["paden"]


def test_boom_slaat_ambigue_titels_over():
    """C# en C++ vallen na het strippen van interpunctie samen; dan liever niets koppelen."""
    catalog = _mini_catalog()
    boom = _mini_boom(catalog, {"name": "wortel", "children": [
        {"name": "Software Development", "children": [
            {"name": "Programmeertalen", "children": [{"name": "Cursus C＃ Professional"}]}]}]})
    # de exacte titel matcht niet en de losse sleutel is ambigu -> geen van beide gekoppeld
    assert "c# professional" not in boom["paden"]
    assert "c++ professional" not in boom["paden"]


def test_boom_kent_meerdere_takken():
    catalog = _mini_catalog()
    boom = _mini_boom(catalog, {"name": "wortel", "children": [
        {"name": "Software Development", "children": [
            {"name": "Web Development", "children": [{"name": "Cursus XSL"}]}]},
        {"name": "Data & Analytics", "children": [
            {"name": "Dataformaten", "children": [{"name": "Cursus XSL"}]}]}]})
    assert len(boom["paden"]["xsl"]) == 2


def test_shortlist_zonder_boom_is_ongewijzigd():
    """De boom is optioneel; zonder boom moet de oude volgorde er exact uit komen."""
    catalog = _mini_catalog()
    zonder = rw.shortlist_vervolgtrainingen(catalog, "Cursus XSL", "stylesheets en xml", 1)
    met_none = rw.shortlist_vervolgtrainingen(catalog, "Cursus XSL", "stylesheets en xml", 1,
                                              boom=None)
    assert [e["titel"] for e in zonder] == [e["titel"] for e in met_none]
    assert "XSL" not in [e["titel"] for e in zonder], "de training zelf hoort er niet in"


def test_shortlist_haalt_vakgenoot_zonder_woordoverlap_binnen():
    """XSL deelt geen woord met JavaScript, maar wel het subdomein Web Development."""
    catalog = _mini_catalog()
    boom = _mini_boom(catalog)
    zonder = {e["titel"] for e in
              rw.shortlist_vervolgtrainingen(catalog, "Cursus XSL", "stylesheets", 1, n=3)}
    met = {e["titel"] for e in
           rw.shortlist_vervolgtrainingen(catalog, "Cursus XSL", "stylesheets", 1, n=3, boom=boom)}
    assert "JavaScript" not in zonder
    assert "JavaScript" in met
    assert "Bloemschikken" not in met


def test_shortlist_houdt_de_sterkste_keyword_treffer_vast():
    """Een vol subdomein mag de beste treffer uit een ánder domein niet verdringen.

    LDAP hangt onder Netwerken, maar de logische vervolgstap (Active Directory) staat
    onder Identity. Die moet de unie overleven.
    """
    catalog = _mini_catalog()
    boom = _mini_boom(catalog)
    titels = [e["titel"] for e in rw.shortlist_vervolgtrainingen(
        catalog, "Cursus LDAP", "directory-gegevens centraal beheren", 4, n=3, boom=boom)]
    assert "Active Directory" in titels


def test_taxonomie_pad_kiest_de_tak_die_de_bron_deelt():
    """Hangt een kandidaat in meerdere takken, dan telt de tak van de bron."""
    catalog = _mini_catalog()
    boom = _mini_boom(catalog, {"name": "wortel", "children": [
        {"name": "Software Development", "children": [
            {"name": "Web Development", "children": [
                {"name": "Cursus XSL"}, {"name": "Cursus JavaScript"}]}]},
        {"name": "Data & Analytics", "children": [
            {"name": "Dataformaten", "children": [{"name": "Cursus JavaScript"}]}]}]})
    assert rw.taxonomie_pad(boom, "javascript", "xsl") == "Software Development > Web Development"
    assert rw.taxonomie_pad(boom, "javascript", "") in (
        "Software Development > Web Development", "Data & Analytics > Dataformaten")
    assert rw.taxonomie_pad(boom, "bloemschikken", "xsl") == ""


class _StubBlok:
    type = "tool_use"

    def __init__(self, naam, invoer):
        self.name, self.input = naam, invoer


class _StubResp:
    def __init__(self, content, stop_reason="tool_use"):
        self.content, self.stop_reason = content, stop_reason


class _StubStroom:
    """Wat `client.messages.stream(...)` teruggeeft: een contextmanager met één bericht.

    `_call_tool` streamt, dus de stubs moeten die vorm hebben. Een stub die `create`
    aanbiedt test een pad dat de pijplijn niet meer loopt -- en juist dat pad heeft een
    plafond op max_tokens dat de streamende variant niet heeft.

    Iterabel, net als het echte object: `_call_tool` loopt de events zelf langs om het
    tijdsbudget tussendoor te kunnen bewaken. Een stub die dat niet kan, test een vorm die
    de SDK niet heeft.
    """

    def __init__(self, resp, events=1):
        self._resp, self._events = resp, events

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def __iter__(self):
        return iter(range(self._events))

    def get_final_message(self):
        return self._resp


class _StubMessages:
    def __init__(self, client, groepen):
        self._client, self._groepen = client, groepen

    def stream(self, **kw):
        self._client.laatste = kw
        return _StubStroom(
            _StubResp([_StubBlok("submit_vervolgstappen", {"groepen": self._groepen})]))


class _StubClient:
    """Client die één vast tool-antwoord teruggeeft; geen netwerk, geen API-key."""

    def __init__(self, groepen):
        self.laatste = None
        self.messages = _StubMessages(self, groepen)


class _StubAfkapper:
    """Kapt de eerste poging af op `max_tokens`; pas de tweede levert het tool-antwoord."""

    def __init__(self):
        self.budgetten = []

    def stream(self, **kw):
        self.budgetten.append(kw["max_tokens"])
        if len(self.budgetten) == 1:
            return _StubStroom(_StubResp([], stop_reason="max_tokens"))
        return _StubStroom(_StubResp([_StubBlok("submit_x", {"klaar": True})]))

    def create(self, **_):
        raise AssertionError("_call_tool mag niet niet-streamend bellen")


# De grens die de SDK zelf rekent: `3600 * max_tokens / 128000` seconden, afgekapt op tien
# minuten. Vanaf 21334 gooit een niet-streamende `messages.create` een ValueError voordat er
# iets over de lijn gaat -- geen API-fout dus, maar een weigering in de client.
NIET_STREAMEND_PLAFOND = 128_000 * 600 // 3600


def test_call_tool_verdubbelt_het_budget_en_doet_dat_streamend():
    """De retry vraagt meer dan een niet-streamende call ooit mag vragen.

    Dit pad was nergens gedekt en viel daarom pas op in productie: de judge maakte zijn
    16000 op, `_call_tool` verdubbelde naar 32000, en de SDK weigerde die tweede call.
    Dat kostte batch 1 in één keer alle 46 trainingen. Twee dingen liggen hier vast: dat
    de verdubbeling gebeurt, en dat ze streamend gebeurt.
    """
    berichten = _StubAfkapper()
    uitkomst = rw._call_tool(SimpleNamespace(messages=berichten), "sys", "user",
                             [{"name": "submit_x"}], "submit_x",
                             max_tokens=16000, thinking=None)
    assert uitkomst == {"klaar": True}
    assert berichten.budgetten == [16000, 32000]
    assert berichten.budgetten[1] > NIET_STREAMEND_PLAFOND


class _StubEindeloos:
    """Een stream die blijft komen, en waarbij de klok tijdens het streamen doorloopt."""

    def __init__(self, events, verstrijk_na):
        self._events, self._verstrijk_na = events, verstrijk_na

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def __iter__(self):
        for i in range(100):
            self._events.append(i)
            if len(self._events) == self._verstrijk_na:
                rw._deadline = time.monotonic() - 1     # de tijd is om, midden in de stream
            yield i

    def get_final_message(self):
        raise AssertionError("deze call had al afgebroken moeten zijn")


def test_call_tool_breekt_een_lopende_call_af_zodra_het_budget_op_is():
    """De grens moet TIJDENS een call gelden, niet pas ertussen.

    Training 47 draaide 81 minuten. De SDK-timeout van 600 s leest als een grens per call maar
    telt per stukje dat over de lijn komt, en een ReadTimeout midden in een stream gaat buiten
    de retry-laag van de SDK om. Een controle tussen de calls door bindt dat dus niet: één
    call kan in zijn eentje langer duren dan het hele budget.
    """
    events = []
    berichten = SimpleNamespace(stream=lambda **kw: _StubEindeloos(events, verstrijk_na=3))
    with rw.tijdsbudget(60):
        try:
            rw._call_tool(SimpleNamespace(messages=berichten), "sys", "user",
                          [{"name": "submit_x"}], "submit_x", thinking=None)
            raise AssertionError("had TijdOverschreden moeten gooien")
        except rw.TijdOverschreden as e:
            assert "tijdsbudget" in str(e), e
    assert len(events) == 3, events      # afgebroken bij het event waarop de tijd om was


def test_een_verstreken_budget_kost_geen_nieuwe_call():
    """Tussen de calls door telt de grens ook, anders begint er nog een dure ronde."""
    gebeld = []
    berichten = SimpleNamespace(stream=lambda **kw: gebeld.append(kw))
    with rw.tijdsbudget(-1):             # deadline in het verleden
        try:
            rw._call_tool(SimpleNamespace(messages=berichten), "sys", "user",
                          [{"name": "submit_x"}], "submit_x", thinking=None)
            raise AssertionError("had TijdOverschreden moeten gooien")
        except rw.TijdOverschreden:
            pass
    assert gebeld == []


def test_tijdsbudget_herstelt_de_vorige_stand_ook_na_een_fout():
    """Een budget dat blijft hangen laat de volgende training op de vorige deadline lopen."""
    assert rw._deadline is None
    with contextlib.suppress(ValueError):
        with rw.tijdsbudget(60):
            assert rw._deadline is not None
            raise ValueError("boem")
    assert rw._deadline is None
    with rw.tijdsbudget(None):           # geen budget -> geen bewaking, geen fout
        rw._bewaak_tijd("wat dan ook")


def test_kies_vervolgtrainingen_toont_het_vakgebied_maar_neemt_het_niet_over():
    """Het label stuurt de groepering; komt het terug in een titel, dan wordt het gestript."""
    catalog = _mini_catalog()
    boom = _mini_boom(catalog)
    shortlist = rw.shortlist_vervolgtrainingen(catalog, "Cursus XSL", "stylesheets", 1, boom=boom)
    client = _StubClient([{"intro": "Verdiep je verder:",
                           "titels": ["JavaScript [Software Development > Web Development]",
                                      "Node.js", "C# Professional"]}])
    groepen = rw.kies_vervolgtrainingen(client, "Training XSL", "stylesheets", "A", shortlist,
                                        boom=boom, oude_titel="Cursus XSL")
    assert "[Software Development > Web Development]" in client.laatste["messages"][0]["content"]
    assert groepen == [{"intro": "Verdiep je verder:",
                        "titels": ["JavaScript", "Node.js", "C# Professional"]}]


def test_kies_vervolgtrainingen_weert_verzonnen_titels():
    catalog = _mini_catalog()
    client = _StubClient([{"intro": "Kijk ook naar:", "titels": ["Training Bestaat Niet"]}])
    shortlist = rw.shortlist_vervolgtrainingen(catalog, "Cursus XSL", "stylesheets", 1)
    assert rw.kies_vervolgtrainingen(client, "Training XSL", "x", "A", shortlist) == []


def test_modules_opening_verdubbelt_soortwoord_niet():
    opening = sjabloon.modules_opening("Opleiding PHP Professional")
    assert "de training PHP Professional" in opening
    assert "Opleiding" not in opening


# ---------------------------------------------------------------------------
# Flags (mogen GEEN hard-fail zijn)
# ---------------------------------------------------------------------------

def test_u_vorm_is_flag_geen_hardfail():
    rw = _good_rewrite()
    rw["doelgroep"] = "Deze training is bedoeld voor iedereen die uw data wil benutten."
    issues = check_rewrite(rw, _CTX)
    assert "u_vorm" in _codes(issues, FLAG)
    assert "u_vorm" not in _codes(issues, HARD)


def test_llm_frase_is_flag():
    rw = _good_rewrite()
    rw["inleiding"] = "In deze training duiken we in " + vul(186, "onderwerp", "thema")
    assert "llm_taal" in _codes(check_rewrite(rw, _CTX), FLAG)


def test_marketing_is_flag():
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je deze uniek " + vul(55) + "?"
    assert "marketing" in _codes(check_rewrite(rw, _CTX), FLAG)


def test_catalogus_niet_geladen_is_flag():
    rw = _good_rewrite()
    issues = check_rewrite(rw, {"naam": "x"})  # geen catalog_titles
    assert "catalogus_ontbreekt" in _codes(issues, FLAG)


# ---------------------------------------------------------------------------
# Tiers: wat komt er in de kolom die een reviewer leest?
#
# Over de eerste 16 herschreven trainingen was 62% van alle flags `lengte_richtlijn` of
# `zwakke_formulering` -- een kolom die daarvoor twee derde uit ruis bestond. De tier zegt
# hoeveel aandacht een flag vraagt; HARD/FLAG blijft zeggen wie hem oplost.
# ---------------------------------------------------------------------------

def _issue(code, section="overzicht", message="boodschap", severity=FLAG):
    return checks.Issue(section, severity, code, message)


def test_elke_tier_code_bestaat_ook_echt_als_check():
    """Een typefout in TIER_PER_CODE laat een flag stilzwijgend op `hoog` staan.

    Dat valt nergens op: de kolom klopt nog, hij is alleen weer even lang als voorheen.
    Vandaar deze bewaker over de broncode -- dezelfde reden als de kolomvolgorde-test.
    """
    bron = open(os.path.join(os.path.dirname(__file__), "rewrite_checks.py")).read()
    bestaande = set(re.findall(r'Issue\([^()]*?,\s*"([a-z_]+)"', bron))
    bestaande |= set(re.findall(r'(?:HARD|FLAG|severity),\s*"([a-z_]+)"', bron))
    onbekend = set(checks.TIER_PER_CODE) - bestaande
    assert not onbekend, f"TIER_PER_CODE noemt codes die geen enkele check maakt: {onbekend}"
    assert set(checks.TIER_PER_CODE.values()) <= set(checks.TIERS)


def test_lengte_binnen_de_vangrail_is_lage_tier():
    """De aanleiding: 13 van de 34 flags waren dit, geen enkele in de buurt van de vangrail."""
    rw = _good_rewrite()
    rw["overzicht"] = "Wil je data kunnen " + vul(82) + "?"     # 86 woorden; band 55-80
    issues = check_rewrite(rw, _CTX)
    assert not hard_fails(issues), [str(i) for i in hard_fails(issues)]
    kolommen = checks.per_tier(flags(issues))
    assert any("86 woorden" in r for r in kolommen[checks.TIER_LAAG])
    assert not any("86 woorden" in r for r in kolommen[checks.TIER_HOOG])


def test_een_oordeel_blijft_hoog_en_een_woordvervanging_wordt_mechanisch():
    kolommen = checks.per_tier([_issue("lerend_aspect"), _issue("anglicisme"),
                                _issue("lengte_richtlijn")])
    assert len(kolommen[checks.TIER_HOOG]) == 1
    assert len(kolommen[checks.TIER_MECHANISCH]) == 1
    assert len(kolommen[checks.TIER_LAAG]) == 1


def test_onbekende_code_valt_op_hoog():
    """De goede kant om op te falen: een nieuwe check komt binnen als werk voor een mens."""
    assert checks.tier(_issue("een_code_die_nog_niet_bestaat")) == checks.TIER_HOOG
    # ook een HARD-issue, dat via `neem_over` in de kolom kan belanden
    assert checks.tier(_issue("modules_te_weinig", severity=HARD)) == checks.TIER_HOOG


def test_dezelfde_opmerking_in_twee_kopjes_wordt_een_regel():
    """Training 27 kreeg "zelfstandig" in het Overzicht en de Inleiding: één beslissing.

    De hoofdletter verschilt omdat de boodschap het gevonden woord citeert en dat aan het
    begin van een zin met een hoofdletter staat; op die twee mag het niet stukgaan.
    """
    kolommen = checks.per_tier([
        _issue("zwakke_formulering", "overzicht", "'zelfstandig': voegt weinig toe."),
        _issue("zwakke_formulering", "inleiding", "'Zelfstandig': voegt weinig toe."),
        _issue("zwakke_formulering", "doelen", "'plaatsen': zegt niet wat de deelnemer kan."),
    ])
    regels = kolommen[checks.TIER_HOOG]
    assert len(regels) == 2, regels
    assert "overzicht + inleiding" in regels[0]
    # de volgorde van de kopjes is die van de checks zelf, niet alfabetisch
    assert regels[0].endswith("'zelfstandig': voegt weinig toe.")


def test_review_tabblad_scheidt_de_drie_tiers():
    """De kolom die naar de reviewer gaat, houdt alleen wat om een oordeel vraagt."""
    res = rw.RewriteResult(1, "Training X", rw.APPROVED, flags_tier={
        checks.TIER_HOOG: ["[FLAG] overzicht: mist het lerende aspect."],
        checks.TIER_MECHANISCH: ["[FLAG] modules: 'Insights' -- schrijf inzichten."],
        checks.TIER_LAAG: ["[FLAG] inleiding: 214 woorden; richtlijn is 180-210 woorden."],
    }, flags=["a", "b", "c"])
    rij = rw._review_rij(res, {})
    assert rij["flags_hoog"] == "[FLAG] overzicht: mist het lerende aspect."
    assert "Insights" in rij["flags_mechanisch"]
    assert "214 woorden" in rij["flags_laag"]
    assert rij["n_hoog"] == 1 and rij["n_flags"] == 3


def test_zonder_tiers_komt_alles_in_de_hoge_kolom():
    """Een oud resultaat of een error-route: liever te veel tonen dan iets verstoppen."""
    res = rw.RewriteResult(1, "Training X", rw.APPROVED, flags=["[FLAG] overzicht: iets."])
    rij = rw._review_rij(res, {})
    assert rij["flags_hoog"] == "[FLAG] overzicht: iets."
    assert rij["flags_mechanisch"] == "" and rij["flags_laag"] == ""


# ---------------------------------------------------------------------------
# Besluiten: structurele splitsing van actie_besluit
#
# De strings hieronder komen letterlijk uit `Nieuwe lijst herschreven en dagen.xlsx`.
# Ze dekken de gevallen waarop de oude regex stukliep: een komma binnen de vrije tekst,
# en een cijfer binnen de vrije tekst ("PHP 8").
# ---------------------------------------------------------------------------

ECHTE_BESLUITEN = [
    ("1,2,3,4", 4),
    ("1,2", 2),
    ("1 niet,2 wel,3 wel,4 wel,5 wel", 5),
    ("1 prima,2 nuanceer,3 niet,4 prima", 4),
    ("1 nee dat is advanced,2 nee,3 in inleiding is dat prima", 3),
    ("1 PHP versie niet benoemen, wel relavante taalfeatures toevoegen,2 prima,3 prima,"
     "4 geen speciefieke frameworks benoemen,5 prima", 5),
    ("1 beide voor zover browser variant nog relevant,2 geen versienummers gebruiken,"
     "3 prima,4 prima", 4),
    ("1 stuk over certificatie vervalt helemaal,2 wel,3 examenstructuur vervalt volledig", 3),
]


def test_besluit_splitst_op_nummer_niet_op_elke_komma():
    for ruw, verwacht in ECHTE_BESLUITEN:
        items = bes.split_besluit(ruw)
        assert len(items) == verwacht, f"{ruw!r} -> {len(items)} i.p.v. {verwacht}"
        assert [nr for nr, _ in items] == list(range(1, verwacht + 1)), ruw


def test_komma_binnen_vrije_tekst_splitst_niet():
    items = bes.split_besluit(
        "1 PHP versie niet benoemen, wel relavante taalfeatures toevoegen,2 prima")
    assert items[0] == (1, "PHP versie niet benoemen, wel relavante taalfeatures toevoegen")
    assert items[1] == (2, "prima")


def test_cijfer_in_vrije_tekst_wordt_geen_actienummer():
    # de oude regex vond hier "8" van "PHP 8" als actienummer
    items = bes.split_besluit("1 gebruik PHP 8 niet als voorbeeld,2 prima")
    assert [nr for nr, _ in items] == [1, 2]


# ---------------------------------------------------------------------------
# Besluiten: dezelfde cel zonder scheidingskomma's (één reviewer doet dat structureel)
# ---------------------------------------------------------------------------

def _acties(n):
    return {i: f"refresh: actie {i}" for i in range(1, n + 1)}


def test_echte_cellen_zonder_scheidingskommas_lijnen_ook_uit():
    """Dezelfde fixtures, met alleen de scheidingskomma's weg -- de rest blijft staan."""
    for ruw, verwacht in ECHTE_BESLUITEN:
        kaal = re.sub(r",\s*(?=\d)", " ", ruw)
        gekoppeld, lezing = bes.koppel_met_lezing(_acties(verwacht), kaal)
        assert [nr for nr, _, _ in gekoppeld] == list(range(1, verwacht + 1)), kaal
        assert lezing == bes.LEZING_NUMMERS, kaal
    # en de vrije tekst blijft heel, inclusief de komma die er wél in hoort
    gekoppeld = bes.koppel(_acties(2), "1 PHP versie niet benoemen, wel taalfeatures 2 prima")
    assert gekoppeld[0][2] == "PHP versie niet benoemen, wel taalfeatures"
    assert gekoppeld[1][2] == "prima"


def test_kommalezing_gaat_voor():
    """Een cel die vandaag werkt wordt niet ineens anders gelezen."""
    ruw = "1 PHP versie niet benoemen, wel taalfeatures,2 prima"
    gekoppeld, lezing = bes.koppel_met_lezing(_acties(2), ruw)
    assert lezing == bes.LEZING_KOMMAS
    assert gekoppeld[0][2] == "PHP versie niet benoemen, wel taalfeatures"


def test_half_gekommade_cel_valt_terug_op_de_nummers():
    gekoppeld = bes.koppel(_acties(3), "1 prima, 2 niet 3 wel")
    assert [ann for _, _, ann in gekoppeld] == ["prima", "niet", "wel"]


def test_cijfer_in_vrije_tekst_kaapt_geen_nummer_zonder_kommas():
    """"versie 3" staat vóór actie 2, dus als scheiding voor 3 valt hij af."""
    gekoppeld = bes.koppel(_acties(3), "1 noem versie 3 niet 2 prima 3 wel")
    assert [ann for _, _, ann in gekoppeld] == ["noem versie 3 niet", "prima", "wel"]


def test_meerduidige_cel_zonder_kommas_is_een_harde_fout():
    """Twee lezingen van hetzelfde nummer -> naar de mens, niet gokken."""
    try:
        bes.koppel(_acties(3), "1 prima 2 noem versie 3 3 wel")
    except bes.BesluitFout as e:
        assert "komma" in str(e)
        return
    raise AssertionError("verwachtte BesluitFout bij een meerduidige cel")


def test_ontbrekend_nummer_noemt_beide_lezingen():
    try:
        bes.koppel(_acties(3), "1 prima 2 niet")
    except bes.BesluitFout as e:
        assert "1 besluiten tegenover 3" in str(e) and "zonder komma's" in str(e)
        return
    raise AssertionError("verwachtte BesluitFout bij een ontbrekend nummer")


# ---------------------------------------------------------------------------
# Scoresheet-kolomnamen: handgemaakte lijsten houden vaak `id`/`name`
# ---------------------------------------------------------------------------

def _scored_df(**kolommen):
    import pandas as pd
    basis = {"actualiteit_actie": ["1. refresh: eerste"], "actie_besluit": ["1 prima"]}
    return pd.DataFrame({**basis, **kolommen})


def test_normaliseer_hernoemt_bronkolomnamen():
    df = bes.normaliseer_scored_kolommen(_scored_df(id=[42], name=["Training XML"]))
    assert "training_id" in df.columns and "titel" in df.columns
    assert df["training_id"].iloc[0] == 42
    assert df["titel"].iloc[0] == "Training XML"


def test_normaliseer_laat_een_scorersheet_met_rust():
    """Staat `training_id` er al, dan wint die -- ook naast een losse `id`-kolom."""
    df = bes.normaliseer_scored_kolommen(
        _scored_df(training_id=[42], titel=["Training XML"], id=[999], name=["iets anders"]))
    assert df["training_id"].iloc[0] == 42
    assert df["titel"].iloc[0] == "Training XML"


def test_load_scored_accepteert_id_en_name():
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "prio.xlsx")
        _scored_df(id=[42], name=["Training XML"]).to_excel(pad, index=False)
        assert bes._load_scored(pad)["training_id"].iloc[0] == 42
        # Via de helper, want dit sheet mist de modus-kolommen en zou hier naar stderr
        # waarschuwen -- terecht, maar niet wat deze test meet.
        assert _laad_scored_met_stderr(pad)[0]["training_id"].iloc[0] == 42


def _laad_scored_met_stderr(pad: str):
    import contextlib
    import io
    buf = io.StringIO()
    with contextlib.redirect_stderr(buf):
        df = rw._load_scored(pad)
    return df, buf.getvalue()


def test_scoresheet_zonder_modus_kolommen_waarschuwt():
    """Het ruwe scoresheet leest zonder foutmelding in, maar draait alles op de defaults.

    Dit is de val die een hele batch kostte: `modus` viel terug op `volledig` en `modules_nb`
    op `stabiel`, terwijl sectie 3b iets anders had voorgesteld. Stil doorgaan mag niet meer.
    """
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "prio.xlsx")
        _scored_df(id=[42], name=["Training XML"]).to_excel(pad, index=False)
        _, melding = _laad_scored_met_stderr(pad)
    assert "modus_voorstel" in melding and "modules_nb_voorstel" in melding
    assert "volledig" in melding and "stabiel" in melding


def test_scoresheet_met_modus_kolommen_zwijgt():
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "met_modus.xlsx")
        _scored_df(id=[42], name=["Training XML"], modus_voorstel=["format"],
                   modules_nb_voorstel=["actueel"], modus_reviewer=[""], kern_reviewer=[""],
                   rewrite_guidance=[""], modules_nb_reviewer=[""]).to_excel(pad, index=False)
        _, melding = _laad_scored_met_stderr(pad)
    assert melding == "", melding


def test_load_scored_faalt_nog_steeds_zonder_id_kolom():
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "zonder_id.xlsx")
        _scored_df(willekeurig=["x"]).to_excel(pad, index=False)
        try:
            bes._load_scored(pad)
        except ValueError as e:
            assert "training_id" in str(e)
            return
    raise AssertionError("verwachtte ValueError bij een sheet zonder id-kolom")


# ---------------------------------------------------------------------------
# Een gedownload blad uit de gedeelde reviewsheet
#
# Het reviewen gebeurt met een team in één Google Sheet. Een blad daaruit downloaden en
# rechtstreeks aan de pijplijn geven moet werken: er wordt op kolomNAAM gematcht, dus de
# volgorde doet niet mee en kolommen die het team zelf bijhoudt rijden ongelezen mee. Deze
# twee tests leggen dat vast, want het is een belofte aan een werkwijze en niet aan één functie.
# ---------------------------------------------------------------------------

def _review_sheet_rijen() -> dict:
    """Eén training met alle kolommen die de pijplijn leest, in dict-vorm."""
    return {
        "training_id": [2347], "titel": ["Cursus Big Data Foundation"],
        "kern": ["Scorer leest de training als introducerend."],
        "kern_reviewer": ["Introducerend: kennismaken met het proces."],
        "actualiteit_actie": ["1. refresh: eerste"], "actie_besluit": ["1 prima"],
        "rewrite_guidance": ["Leg de nadruk op governance."],
        "verdict": ["redelijk"], "herschreven": [0],
        "vermoedelijk_persona": ["A"], "aantal_dagen_bron": [2],
        "actualiteit_type": ["additief"],
        "bruikbaar": ["Praktijkcase | Module een"], "strippen": ["Alternatief-paragraaf"],
        "gaten": ["Geen doelgroep"], "menselijke_input_nodig": [False],
        "modus_voorstel": ["format"], "modules_nb_voorstel": ["stabiel"],
        "modus_reviewer": [""], "modules_nb_reviewer": [""],
    }


def _schrijf_bron(pad: str) -> None:
    import pandas as pd
    pd.DataFrame({"id": [2347], "name": ["Cursus Big Data Foundation"],
                  "content": [json.dumps({"days": 2, "intro": "<p>CRISP-DM.</p>"})]}
                 ).to_excel(pad, index=False)


def test_gedownload_reviewblad_levert_dezelfde_briefing():
    """Kolomvolgorde en vreemde kolommen veranderen de briefing niet -- alles matcht op naam."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        bron = os.path.join(d, "bron.xlsx")
        kaal = os.path.join(d, "kaal.xlsx")
        uit_de_sheet = os.path.join(d, "uit_de_sheet.xlsx")
        _schrijf_bron(bron)

        rijen = _review_sheet_rijen()
        pd.DataFrame(rijen).to_excel(kaal, index=False)

        # zoals het blad er na een teamronde uitziet: omgekeerde volgorde, plus de kolommen die
        # het team zelf bijhoudt -- inclusief een formule, want zo staat de CRM-link erin
        gehusseld = {k: rijen[k] for k in reversed(list(rijen))}
        gehusseld.update({"Verantwoordelijk voor input": ["Remko"], "Status": ["nagekeken"],
                          "Link naar CRM": ['=HYPERLINK("https://crm.eduvision.nl/", "crm")'],
                          "Link naar herschreven training": [None]})
        pd.DataFrame(gehusseld).to_excel(uit_de_sheet, index=False)

        a = rw.build_briefing_for_id(kaal, bron, 2347)
        b = rw.build_briefing_for_id(uit_de_sheet, bron, 2347)

    assert a == b, "de gehusselde, verrijkte sheet levert een andere briefing"
    # niet alleen gelijk, ook goed: de velden die de reviewer vulde zijn echt aangekomen
    assert b.modus == "format", b.modus
    assert b.kern_definitief.startswith("Introducerend"), b.kern_definitief
    assert "governance" in b.guidance_definitief


def test_modus_voorstellen_houdt_de_kolomvolgorde_van_het_reviewsheet():
    """Het modus-blok komt achteraan; de scorer-kolommen houden de volgorde van het sheet.

    `met_llm=False` doet geen enkele API-call: `schat_modus` valt bij `client is None` meteen
    terug op de deterministische ondergrens.
    """
    import pandas as pd
    import score_trainings as st
    with tempfile.TemporaryDirectory() as d:
        bron = os.path.join(d, "bron.xlsx")
        scored = os.path.join(d, "scored.xlsx")
        uit_pad = os.path.join(d, "met_modus.xlsx")
        _schrijf_bron(bron)

        rijen = {k: v for k, v in _review_sheet_rijen().items()
                 if k not in ("modus_voorstel", "modules_nb_voorstel",
                              "modus_reviewer", "modules_nb_reviewer")}
        rijen["Status"] = ["nagekeken"]
        pd.DataFrame(rijen).to_excel(scored, index=False)

        rw.modus_voorstellen(scored, bron, uit_pad, met_llm=False, verbose=False)
        kolommen = list(pd.read_excel(uit_pad).columns)

    canoniek = [k for k in st.KOLOM_VOLGORDE if k in kolommen]
    assert kolommen[:len(canoniek)] == canoniek, f"canonieke voorloop gebroken: {kolommen}"
    for modus_kolom in ("modus_voorstel", "modus_reden", "modus_ondergrens",
                        "modules_nb_voorstel", "modules_nb_reden", "modus_reviewer",
                        "modules_nb_reviewer"):
        assert kolommen.index(modus_kolom) >= len(canoniek), \
            f"{modus_kolom} staat binnen het plakblok"
    # de kolom van het reviewteam overleeft, en `guidance_reviewer` komt er niet meer bij
    assert "Status" in kolommen
    assert "guidance_reviewer" not in kolommen, "de legacy-kolom wordt weer aangemaakt"


# ---------------------------------------------------------------------------
# De id-kolom: Excel maakt er zomaar floats van
# ---------------------------------------------------------------------------

def test_hele_floats_worden_weer_gehele_ids():
    """`796.0` in beeld en `796.000` in het sheet is voor een reviewer niet over te typen."""
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "prio.xlsx")
        _scored_df(id=[796.0], name=["Training XML"]).to_excel(pad, index=False)
        import contextlib
        import io
        for laag in (bes, rw):
            # stderr weg: zie de toelichting bij de duizendtal-test hieronder.
            with contextlib.redirect_stderr(io.StringIO()):
                tid = laag._load_scored(pad)["training_id"].iloc[0]
            assert tid == 796 and int(tid) == tid, (laag.__name__, tid)
            assert "." not in str(tid), (laag.__name__, tid)


def test_id_met_duizendtalscheiding_wordt_geweigerd():
    """`2.347` is 2347 met een punt erin; als decimaal gelezen joint hij nergens meer mee."""
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "prio.xlsx")
        _scored_df(id=[2.347], name=["Cursus Big Data Foundation"]).to_excel(pad, index=False)
        import contextlib
        import io
        for laag in (bes, rw):
            try:
                # stderr weg: `rw._load_scored` waarschuwt hier terecht over de ontbrekende
                # modus-kolommen, maar dat is niet wat deze test meet.
                with contextlib.redirect_stderr(io.StringIO()):
                    laag._load_scored(pad)
            except ValueError as e:
                assert "2.347" in str(e) and "duizendtal" in str(e), str(e)
            else:
                raise AssertionError(f"{laag.__name__} accepteerde een niet-heel training_id")


def test_tekstuele_ids_blijven_ongemoeid():
    """Niet elk id is een getal; alleen de float-val hoort hier te worden afgevangen."""
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "prio.xlsx")
        _scored_df(id=["XML-42"], name=["Training XML"]).to_excel(pad, index=False)
        assert _laad_scored_met_stderr(pad)[0]["training_id"].iloc[0] == "XML-42"


# ---------------------------------------------------------------------------
# besluiten.xlsx: meerdere batches naast elkaar
# ---------------------------------------------------------------------------

def test_besluiten_sheet_behoudt_trainingen_buiten_het_scoresheet():
    """Een prioriteitslijst mag de besluiten van de rest van de catalogus niet wissen."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        besluiten_pad = os.path.join(d, "besluiten.xlsx")
        eerste = os.path.join(d, "batch_a.xlsx")
        tweede = os.path.join(d, "batch_b.xlsx")
        _scored_df(training_id=[1], titel=["Training A"]).to_excel(eerste, index=False)
        _scored_df(id=[2], name=["Training B"]).to_excel(tweede, index=False)

        bes.write_besluiten_sheet(eerste, besluiten_pad, verbose=False)
        bes.write_besluiten_sheet(tweede, besluiten_pad, verbose=False)

        df = pd.read_excel(besluiten_pad)
        assert set(df["training_id"]) == {1, 2}
        assert list(df.columns) == bes.KOLOMMEN
        assert set(bes.load_besluiten(besluiten_pad)) == {1, 2}


def test_besluiten_sheet_ververst_de_eigen_trainingen():
    """Opnieuw draaien op hetzelfde sheet dupliceert niet en neemt de nieuwe actietekst mee."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        besluiten_pad = os.path.join(d, "besluiten.xlsx")
        scored = os.path.join(d, "batch.xlsx")
        _scored_df(training_id=[1], titel=["Training A"]).to_excel(scored, index=False)
        bes.write_besluiten_sheet(scored, besluiten_pad, verbose=False)

        gewijzigd = _scored_df(training_id=[1], titel=["Training A"])
        gewijzigd["actualiteit_actie"] = ["1. refresh: herzien"]
        gewijzigd.to_excel(scored, index=False)
        bes.write_besluiten_sheet(scored, besluiten_pad, verbose=False)

        df = pd.read_excel(besluiten_pad)
        assert len(df) == 1
        assert df["actie"].iloc[0] == "refresh: herzien"


def test_parse_acties_nummert_de_scorerlijst():
    acties = bes.parse_acties("1. refresh: eerste\n2. refresh: tweede\n3. refresh: derde")
    assert acties == {1: "refresh: eerste", 2: "refresh: tweede", 3: "refresh: derde"}


def test_align_koppelt_annotatie_aan_actie():
    acties = {1: "eerste", 2: "tweede"}
    gekoppeld = bes.align(acties, bes.split_besluit("1 niet,2 prima"))
    assert gekoppeld == [(1, "eerste", "niet"), (2, "tweede", "prima")]


def test_align_faalt_hard_bij_verkeerd_aantal():
    try:
        bes.align({1: "eerste", 2: "tweede"}, bes.split_besluit("1,2,3"))
    except bes.BesluitFout:
        return
    raise AssertionError("verwachtte BesluitFout bij 3 besluiten op 2 acties")


def test_align_faalt_hard_bij_onbekend_nummer():
    try:
        bes.align({1: "eerste"}, bes.split_besluit("7 prima"))
    except bes.BesluitFout:
        return
    raise AssertionError("verwachtte BesluitFout bij een nummer zonder actie")


# ---------------------------------------------------------------------------
# Besluiten: de regel-fastpath (exact-match, buiten het model om)
# ---------------------------------------------------------------------------

def test_fastpath_keurt_kale_goedkeuring_goed():
    for ann in ("", "prima", "wel", "ja", "OK", "Akkoord", "prima."):
        assert bes.regel_label(ann) == bes.DOEN, ann


def test_fastpath_wijst_kale_afwijzing_af():
    for ann in ("niet", "nee", "Nee", "nvt"):
        assert bes.regel_label(ann) == bes.NIET, ann


def test_fastpath_laat_vrije_tekst_aan_het_model():
    # "geen ..." is een VOORWAARDE, geen afwijzing -- de fastpath mag hier niet gokken
    for ann in ("geen specifieke frameworks benoemen", "nee dat is advanced",
                "in inleiding is dat prima", "nuanceer"):
        assert bes.regel_label(ann) is None, ann


# ---------------------------------------------------------------------------
# Besluiten: wat de schrijver te zien krijgt
# ---------------------------------------------------------------------------

def _besluit(besluit, voorwaarde=""):
    return bes.Besluit(1, "T", 1, "refresh: doe iets", voorwaarde, besluit,
                       voorwaarde, bes.BRON_LLM)


def test_voorwaarde_gaat_altijd_mee_naar_de_schrijver():
    # ook bij `doen`: het label bepaalt of de actie doorgaat, niet of de reviewer
    # gehoord wordt -- anders verdampt de aantekening bij een net-verkeerd label
    for label, kop in ((bes.MITS, "VOORWAARDE"), (bes.DOEN, "VOORWAARDE"), (bes.NIET, "REDEN")):
        tekst = _besluit(label, "alleen in de inleiding").als_instructie()
        assert "alleen in de inleiding" in tekst
        assert kop in tekst


def test_kale_goedkeuring_voegt_niets_toe():
    assert _besluit(bes.DOEN).als_instructie() == "refresh: doe iets"


def test_splits_scheidt_goedgekeurd_van_afgewezen():
    goed, afgewezen = bes.splits([_besluit(bes.DOEN), _besluit(bes.MITS), _besluit(bes.NIET)])
    assert len(goed) == 2 and len(afgewezen) == 1


# ---------------------------------------------------------------------------
# De kern: het enige veld dat het niveau van de training vastlegt
#
# De reviewer stuurt bij in `kern_reviewer`; de scorer-kern blijft ernaast staan zodat een
# herscoring hem mag verversen. Wie de kern schreef bepaalt hoeveel gezag hij heeft.
# ---------------------------------------------------------------------------

def _briefing(**overrides) -> rw.RewriteBriefing:
    basis = dict(training_id=1, titel="Cursus Big Data Foundation", persona="A", dagen=2,
                 kern="Scorer leest de training als een gestructureerde aanpak.",
                 verdict="redelijk", actualiteit_type="additief",
                 source_text="[intro]\nWe introduceren het CRISP-DM model.")
    return rw.RewriteBriefing(**{**basis, **overrides})


def test_reviewer_kern_wint_van_scorer_kern():
    b = _briefing(kern_reviewer="Introducerend: de deelnemer maakt kennis met het proces.")
    assert b.kern_definitief == "Introducerend: de deelnemer maakt kennis met het proces."
    assert b.kern_van_reviewer


def test_zonder_reviewer_kern_geldt_de_scorer_kern():
    b = _briefing()
    assert b.kern_definitief == b.kern
    assert not b.kern_van_reviewer


def test_lege_reviewer_kern_telt_niet_als_oordeel():
    """Een lege Excel-cel komt binnen als NaN of als spaties; geen van beide is een oordeel."""
    for leeg in ("", "   ", "\n"):
        assert not _briefing(kern_reviewer=leeg).kern_van_reviewer
    assert rw._cel(float("nan")) == ""
    assert rw._cel(None) == ""
    assert rw._cel("  echte tekst  ") == "echte tekst"


def test_gezagsregel_volgt_de_herkomst_van_de_kern():
    """Bij een scorer-kern wint de bron; bij een reviewer-kern wint de kern."""
    scorer = rw.build_writer_user(_briefing())
    assert "lezing van de scorer" in scorer
    assert "wint de BRONTEKST" in scorer

    reviewer = rw.build_writer_user(_briefing(kern_reviewer="Introducerende training."))
    assert "vastgesteld door reviewer" in reviewer
    assert "Hij is leidend" in reviewer
    assert "wint de BRONTEKST" not in reviewer


def test_writer_prompt_zet_de_brontekst_neer_als_de_training_zelf():
    tekst = rw.build_writer_user(_briefing())
    assert "BRONTEKST" in tekst
    assert "Beloof nooit meer dan hier staat." in tekst
    assert tekst.rstrip().endswith("We introduceren het CRISP-DM model.")


def test_judge_krijgt_de_kern_met_herkomst():
    """Zonder de kern kan de judge het niveau niet toetsen -- dat kon hij eerder niet."""
    doc = {"titel": "Training Big Data Foundation", "overzicht": "Wil je iets?"}
    scorer = rw.build_judge_user(_briefing(), doc)
    assert "KERN (lezing van de scorer)" in scorer
    reviewer = rw.build_judge_user(_briefing(kern_reviewer="Introducerend."), doc)
    assert "KERN (vastgesteld door reviewer)" in reviewer
    assert "Introducerend." in reviewer


def test_judge_krijgt_de_brontekst_en_alle_feiten():
    """§2 van de beoordelingsspec vraagt te herleiden tot "de brontekst" -- die kreeg hij niet.

    Zonder de bron toetste de judge feitgetrouwheid tegen de samenvatting van de scorer: een
    claim die de brontekst tegensprak maar `bruikbaar` niet, kwam er ongehinderd doorheen. Dat
    weegt zwaarder sinds de kern het niveau draagt: zwijgt de kern over een aspect, dan had de
    judge niets om op terug te vallen.
    """
    doc = {"titel": "Training Big Data Foundation", "overzicht": "Wil je iets?"}
    tekst = rw.build_judge_user(
        _briefing(bruikbaar=["CRISP-DM 6 fasen"], strippen=["verouderde prijsinfo"],
                  gaten=["voorkennis niet beschreven"]),
        doc)
    assert "We introduceren het CRISP-DM model." in tekst
    # strippen en gaten stonden wél in de spec beloofd, maar gingen niet mee
    assert "verouderde prijsinfo" in tekst
    assert "voorkennis niet beschreven" in tekst
    # het concept staat achteraan: de bron is naslag, het concept is wat hij beoordeelt
    assert tekst.index("BRONTEKST") < tekst.index("CONCEPT. Dit is wat je beoordeelt")


def test_judge_mag_de_bron_niet_als_vormnorm_gebruiken():
    """Zonder deze regel rekent de judge het concept af op het niet volgen van de bronstructuur
    -- en juist daarvan afwijken is het hele punt van herschrijven."""
    tekst = rw.build_judge_user(_briefing(), {"titel": "T", "overzicht": "Wil je iets?"})
    assert "Reken het concept NIET af op vorm" in tekst
    assert "FEITGETROUWHEID" in tekst and "NIVEAU" in tekst


def test_goedgekeurde_actualisering_gaat_voor_de_brontekst():
    """De bron is de maatstaf voor claims -- behalve voor wat de reviewer heeft goedgekeurd.

    Een actualisering voegt per definitie iets toe dat niet in de bron staat of haalt iets weg; dat is waarom
    hij bestaat. Zonder deze uitzondering is elke goedgekeurde actie een "verzonnen feit" en
    draait de review precies het werk terug dat de reviewer in de sessie deed. Geldt aan
    beide kanten: de schrijver mag er geen laten liggen, de judge mag er geen afkeuren.
    """
    actie = bes.Besluit(1, "T", 1, "Voeg GA4 toe (bron noemt alleen Universal Analytics)",
                        "prima, mits als voorbeeld", "mits", "alleen als voorbeeld",
                        "handmatig")
    b = _briefing(goedgekeurd=[actie])

    schrijver = rw.build_writer_user(b)
    assert "Eén uitzondering, en die gaat vóór" in schrijver
    # de uitzondering staat ná "Beloof nooit meer dan hier staat", anders leest de schrijver
    # dat verbod als het laatste woord en laat hij de actie liggen
    assert schrijver.index("Beloof nooit meer dan hier staat") < schrijver.index("uitzondering")

    judge = rw.build_judge_user(b, {"titel": "T", "overzicht": "Wil je iets?"})
    assert "UITZONDERING op punt 1, en die gaat vóór" in judge
    assert "nooit af als ongegrond" in judge
    # ... maar die uitzondering dekt het onderwerp en niet het niveau. Zonder die grens liet
    # de judge training 27 door: "benoem concrete SQL-platformen" werd "pas je direct toe op",
    # en de spec verbood hem letterlijk om dat als te hoge belofte af te rekenen.
    assert "dekt het ONDERWERP, niet het NIVEAU" in judge
    assert "Punt 2 blijft dus gewoon gelden" in judge
    # de voorwaarde is de tweede grens, en die gaat mee naar allebei
    assert "alleen als voorbeeld" in schrijver and "alleen als voorbeeld" in judge


def test_werkwoord_van_de_actie_is_de_bovengrens_in_elke_actie_prompt():
    """Uit training 27: "benoem concrete SQL-platformen" werd "pas je direct toe op".

    De regel hoort in het actualiseringenblok en niet in het modusblok. De enige rem die er
    was, `ACTUALISEREN_ONGEACHT_MODUS`, wordt alleen gerenderd als er een MODUS_UITLEG is; in
    `volledig` kreeg de schrijver dus niets. Vandaar dat deze test juist die modus pakt.
    """
    actie = bes.Besluit(1, "T", 1, "refresh: benoem concrete SQL-platformen als context",
                        "in inleiding is dat prima", "mits", "in inleiding is dat prima",
                        "handmatig")
    b = _briefing(modus_reviewer="volledig", goedgekeurd=[actie])

    schrijver = rw.build_writer_user(b)
    assert "staan LOS van deze opdracht" not in schrijver     # geen modusblok in `volledig`
    for prompt in (schrijver,
                   rw.build_judge_user(b, {"titel": "T", "overzicht": "Wil je iets?"}),
                   rw.build_actualisatie_user(b, _content(), "Training SQL")):
        assert "Het werkwoord van een actie is de bovengrens" in prompt
        assert "Wat je noemt maar niet traint, beloof je niet." in prompt


def test_judge_ziet_de_reden_bij_een_afgewezen_actie():
    """De schrijver kreeg de reden al, de judge niet: die zag alleen `actie`.

    Zonder de reden ziet hij dat iets niet mag, maar niet waaróm, en dat is het verschil
    tussen "staat er niet in" en "is bewust weggehouden".
    """
    afgewezen = bes.Besluit(1, "T", 1, "refresh: voeg window functions toe",
                            "nee dat is advanced", "niet", "nee dat is advanced", "handmatig")
    b = _briefing(afgewezen=[afgewezen])
    judge = rw.build_judge_user(b, {"titel": "T", "overzicht": "Wil je iets?"})
    assert "REDEN (reviewer): nee dat is advanced" in judge


def test_noem_actie_die_toepassen_wordt_is_een_flag():
    """De letterlijke zin uit training 27, met de letterlijke actie ernaast."""
    ctx = {"naam": "Training SQL",
           "acties": ["refresh: benoem concrete SQL-platformen (bv. PostgreSQL, SQL Server, "
                      "cloud data warehouses) als context bij de training"]}
    rwin = {"inleiding": "De training is praktijkgericht opgezet. De SQL die je leert, pas je "
                         "direct toe op verschillende platformen, van PostgreSQL en SQL Server "
                         "tot cloud data warehouses."}
    issues = checks.check_actie_escalatie(rwin, ctx)
    assert len(issues) == 1, issues
    assert issues[0].code == "actie_escalatie" and issues[0].severity == checks.FLAG
    assert issues[0].section == "inleiding"
    # nooit hard: de grens tussen noemen en behandelen is niet deterministisch vast te stellen
    assert not checks.hard_fails(checks.check_actie_escalatie(rwin, ctx))


def test_noem_actie_zonder_escalatie_vuurt_niet():
    """Drie vrijwaringen die de meting over `herschreven/trainingen/` heeft opgeleverd."""
    actie = ["refresh: benoem concrete SQL-platformen (bv. PostgreSQL, SQL Server) als context"]

    # 1. dezelfde platformen, netjes benoemd in plaats van beloofd
    goed = {"inleiding": "De SQL die je leert, werkt op de platformen die je in de praktijk "
                         "tegenkomt, zoals PostgreSQL en SQL Server."}
    assert not checks.check_actie_escalatie(goed, {"naam": "Training SQL", "acties": actie})

    # 2. "van toepassing" is een idioom, geen leeractiviteit (viel op 3077)
    idioom = {"inleiding": "Sinds januari 2025 is PostgreSQL 17 volledig van toepassing binnen "
                           "de standaard."}
    assert not checks.check_actie_escalatie(idioom, {"naam": "Training SQL", "acties": actie})

    # 3. een term die het onderwerp van de training zélf is (viel op 2808)
    eigen = {"inleiding": "Je werkt met PostgreSQL als centrale database."}
    assert not checks.check_actie_escalatie(
        eigen, {"naam": "Training PostgreSQL", "acties": actie})

    # en een actie die wél om behandelen vraagt, mag een leeractiviteit worden
    zwaar = ["refresh: voeg een module toe over PostgreSQL en SQL Server"]
    beloofd = {"inleiding": "Je werkt met PostgreSQL en past SQL Server toe op je eigen queries."}
    assert not checks.check_actie_escalatie(beloofd, {"naam": "Training SQL", "acties": zwaar})


def test_check_ctx_is_voor_beide_aanroepers_hetzelfde():
    """`rewrite_one` en `hergenereer_kopje` bouwden dit allebei zelf; toen `acties` erbij kwam
    was dat meteen een plek waar de ene een check kon draaien die de andere niet had."""
    actie = bes.Besluit(1, "T", 1, "refresh: benoem PostgreSQL", "ja", "doen", "", "handmatig")
    ctx = rw.build_check_ctx(_briefing(goedgekeurd=[actie]), None)
    assert ctx["acties"] == ["refresh: benoem PostgreSQL"]     # kaal, zonder voorwaarde
    assert set(ctx) == {"catalog_titles", "naam", "dagen", "acties"}
    assert "ctx = build_check_ctx(b, catalog)" in open(rw.__file__, encoding="utf-8").read()


def test_judge_oordeel_met_verkeerd_gevormd_blok_loopt_niet_stuk():
    """Het tool-schema dwingt de vorm niet af: één meetrun leverde `feitgetrouw` als string.

    Ongefilterd loopt dat verderop stuk op `judgment["feitgetrouw"].get("thin")` -- midden in
    een batch, ná de dure schrijfcall.
    """
    kapot = {"verdict": "approved", "feitgetrouw": "alles klopt", "judge_confidence": "high"}
    rw._call_tool, echt = (lambda *a, **k: kapot), rw._call_tool
    try:
        out = rw.judge_document(None, _briefing(), {"titel": "T"})
    finally:
        rw._call_tool = echt
    assert out["verdict"] == "approved"          # de rest van het oordeel blijft bruikbaar
    assert out["feitgetrouw"] == {}
    assert "feitgetrouw" in out["judge_vorm"]    # maar het misgaan blijft zichtbaar
    assert out.get("feitgetrouw", {}).get("thin", False) is False


def test_rondes_leggen_per_poging_vast_wat_er_gebeurde():
    """`judgment` bewaart alleen het LAATSTE oordeel, en daarop is `MAX_REVISIONS` niet bij te
    stellen: drie keer dezelfde klacht betekent dat een ronde erbij helpt, elke ronde een
    andere betekent dat de lus niet convergeert. Over batch 1 was dat verschil achteraf niet
    meer te zien -- vijf trainingen liepen op de limiet en niemand kon zeggen waarom.
    """
    catalog = [{"product_id": 9, "titel": "Training Power BI", "summary": ""}]
    klachten = iter(range(10))
    echt = rw._call_tool, rw.bepaal_vervolgstappen, rw.judge_document
    rw._call_tool = lambda *a, **k: _good_rewrite()
    rw.bepaal_vervolgstappen = lambda *a, **k: (["Training Power BI"], [])
    rw.judge_document = lambda *a, **k: {
        "verdict": rw.NEEDS_REVISION, "revisie_notities": [f"klacht {next(klachten)}"]}
    try:
        res = rw.rewrite_one(None, _briefing(titel="Cursus Data-analyse"), catalog)
    finally:
        rw._call_tool, rw.bepaal_vervolgstappen, rw.judge_document = echt

    assert res.status == rw.HUMAN_QUEUE
    assert [r["ronde"] for r in res.rondes] == list(range(1, rw.MAX_REVISIONS + 2))
    assert {r["uitkomst"] for r in res.rondes} == {rw.NEEDS_REVISION}
    # elke ronde zijn eigen notities: precies het spoor dat "dezelfde klacht of een nieuwe?"
    # beantwoordt
    assert res.rondes[0]["notities"] == ["klacht 0"]
    assert res.rondes[-1]["notities"] == [f"klacht {rw.MAX_REVISIONS}"]


def test_build_briefing_leest_de_reviewerkolom():
    scored = {"training_id": 7, "kern": "scorer-kern", "kern_reviewer": "reviewer-kern",
              "verdict": "rijk", "vermoedelijk_persona": "B"}
    b = rw.build_briefing(scored, {"days": 3, "intro": "tekst"}, "Cursus X")
    assert b.kern_definitief == "reviewer-kern" and b.kern_van_reviewer
    # ontbreekt de kolom (oud scoresheet), dan valt hij stil terug op de scorer-kern
    zonder = rw.build_briefing({k: v for k, v in scored.items() if k != "kern_reviewer"},
                               {"days": 3, "intro": "tekst"}, "Cursus X")
    assert zonder.kern_definitief == "scorer-kern" and not zonder.kern_van_reviewer


def test_reviewtabblad_zet_de_brontekst_naast_de_nieuwe_tekst():
    """De mens kon alleen de nieuwe tekst lezen, dus niet zien of een claim de bron dekt."""
    res = rw.RewriteResult(1, "Training Big Data Foundation", rw.APPROVED,
                           oude_titel="Cursus Big Data Foundation")
    rij = rw._review_rij(res, {}, {"intro": "We introduceren het CRISP-DM model."})
    assert "We introduceren het CRISP-DM model." in rij["brontekst"]
    # de kolom staat achteraan, naast de kopjes -- niet tussen de statusvelden
    assert list(rij)[-1] == "brontekst"
    # zonder bron (overgenomen zonder bronrij) blijft de cel gewoon leeg
    assert rw._review_rij(res, {})["brontekst"] == ""


# ---------------------------------------------------------------------------
# Output: document -> CMS-content
# ---------------------------------------------------------------------------

def _document() -> dict:
    return {
        "titel": "Training XML",
        "overzicht": "Wil je slimmer werken?",
        "inleiding": "Eerste alinea.\n\nTweede alinea.",
        "modules": {"opening": sjabloon.modules_opening("Cursus XML"),
                    "modules": [{"titel": "M1", "bullets": ["a", "b"]}]},
        "doelgroep": "Deze training is bedoeld voor iedereen.",
        "voorkennis": sjabloon.VOORKENNIS_FALLBACK,
        # precies zoals `assemble_document` hem samenstelt
        "aanpak": sjabloon.AANPAK_ALINEA_1.format(invulling="je dit toepast")
                  + "\n\n" + sjabloon.AANPAK_ALINEA_2,
        "doelen": {"intro": sjabloon.DOELEN_INTRO, "bullets": ["Doen van dingen"]},
        "vervolgstappen": {"alineas": [sjabloon.VERVOLG_ALINEA_1, sjabloon.VERVOLG_ALINEA_2],
                           "titels": ["Power BI"], "afsluiter": sjabloon.VERVOLG_AFSLUITER},
        "kortste_omschrijving": "Wil je dit leren?",
        "certificatie": sjabloon.CERTIFICATIE,
    }


def test_content_heeft_dezelfde_sleutels_als_de_bron():
    bron = {"days": 3, "intro": "", "setup": "", "modules": "", "summary": "",
            "follow_up": "", "objectives": "", "certification": "", "summary_edudex": "",
            "prior_knowledge": "", "target_audience": ""}
    content = uit.document_to_content(_document(), bron)
    assert set(content) == set(bron), set(bron) ^ set(content)


def test_days_wordt_ongewijzigd_overgenomen():
    content = uit.document_to_content(_document(), {"days": 7})
    assert content["days"] == 7


def test_summary_is_platte_tekst_geen_html():
    content = uit.document_to_content(_document(), {})
    assert "<" not in content["summary"] and "<" not in content["summary_edudex"]


def test_inleiding_krijgt_het_bedrijfstrainingblok_als_kop_3():
    content = uit.document_to_content(_document(), {})
    assert f"<h3>{sjabloon.BEDRIJFSTRAINING_KOP}</h3>" in content["intro"]
    assert sjabloon.BEDRIJFSTRAINING_TEKST in content["intro"]


def test_modules_gebruikt_geneste_lijsten():
    content = uit.document_to_content(_document(), {})
    assert content["modules"].count("<ul>") == 2   # buitenste + één module
    assert "<h3>" not in content["modules"]


def test_geen_placeholder_of_oplnaam_in_de_output():
    content = uit.document_to_content(_document(), {})
    samen = " ".join(v for v in content.values() if isinstance(v, str))
    assert "{{ oplnaam }}" not in samen and "[" not in samen


def test_vervolgstappen_met_groepen_krijgt_per_groep_een_eigen_intro():
    doc = _document()
    doc["vervolgstappen"]["groepen"] = [
        {"intro": "Wil je je verder verdiepen:",
         "titels": ["Training Power BI", "Training DAX"]},
        {"intro": "Wil je juist verbreden:", "titels": ["Training T-SQL", "Training Python"]},
    ]
    html = uit.document_to_content(doc, {})["follow_up"]
    assert "Wil je je verder verdiepen:" in html and "Wil je juist verbreden:" in html
    assert html.count("<ul>") == 2
    # de vaste aankondiging hoort er dan NIET meer boven te staan
    assert sjabloon.VERVOLG_LIJST_INTRO not in html


def test_vervolgstappen_zonder_groepen_valt_terug_op_een_vlakke_lijst():
    html = uit.document_to_content(_document(), {})["follow_up"]
    assert sjabloon.VERVOLG_LIJST_INTRO in html
    assert html.count("<ul>") == 1


def test_soortwoord_wordt_niet_verdubbeld():
    assert "de training opleiding" not in sjabloon.modules_opening("Opleiding PHP Professional").lower()
    assert sjabloon.modules_opening("Cursus XML").startswith("Tijdens de training XML")
    assert sjabloon.modules_opening("Photoshop").startswith("Tijdens de training Photoshop")
    assert sjabloon.modules_opening("Masterclass C#").startswith("Tijdens de masterclass C#")


def test_markdown_heeft_kop_1_2_en_3():
    md = uit.render_markdown(_document())
    assert md.startswith("# Training XML")
    for kopje in sjabloon.KOPJES:
        assert f"## {kopje.kop}" in md, kopje.kop
    assert f"### **{sjabloon.BEDRIJFSTRAINING_KOP}**" in md


def test_de_titel_uit_het_document_wint_van_het_argument():
    """De judge las de mechanische titel in plaats van de gekozen titel, en flagde die.

    `build_judge_user` gaf `b.nieuwe_titel` mee, dus wat `bepaal_titel` had gekozen kwam nooit
    bij de judge aan. Training 279 leverde de goedgekeurde rename ("Training HTML en CSS"),
    had die in zijn document staan, en kreeg drie rondes lang de opdracht om een titel te
    veranderen die al veranderd wás -- een revisielus die niet te winnen is en dus altijd in
    de menselijke wachtrij eindigt.
    """
    doc = dict(_document(), titel="Training HTML en CSS")
    assert uit.render_markdown(doc, "Training HTML5 en CSS3").startswith("# Training HTML en CSS")
    # het argument blijft de terugval voor een document zonder eigen titel
    assert uit.render_markdown({k: v for k, v in doc.items() if k != "titel"},
                               "Training XML").startswith("# Training XML")


def test_de_judge_beoordeelt_de_titel_die_bepaal_titel_koos():
    """Zelfde regel, maar dan op de plek waar hij misging: het concept dat de judge leest."""
    b = _briefing(titel="Cursus HTML5 en CSS3")
    doc = dict(_document(), titel="Training HTML en CSS")
    tekst = rw.build_judge_user(b, doc)
    concept = tekst.split("CONCEPT.")[1]
    assert "# Training HTML en CSS" in concept
    assert "# Training HTML5 en CSS3" not in concept


# ---------------------------------------------------------------------------
# Document -> HTML voor Google Docs
# ---------------------------------------------------------------------------

def _koppen(doc: str, niveau: str) -> list[str]:
    """De kopteksten van dit niveau, zonder de opmaak eromheen.

    Toetsen op `<h2>Doelen</h2>` zou hier meten wat de opmaak doet in plaats van wat er staat;
    de koppen dragen sinds de eerste echte batch een `style` en een `<span>`.
    """
    binnenkant = re.findall(rf"<{niveau}\b[^>]*>(.*?)</{niveau}>", doc, re.S)
    return [re.sub(r"<[^>]+>", "", k) for k in binnenkant]


def test_docs_html_heeft_een_h1_en_tien_h2_koppen():
    """De koppen zijn de winst: Drive maakt er de documentoverzicht-zijbalk van."""
    doc = uit.render_docs_html(uit.document_to_content(_document(), {}), "Cursus XML")
    assert _koppen(doc, "h1") == ["Cursus XML"]
    assert _koppen(doc, "h2") == [k.kop for k in sjabloon.KOPJES]
    # zelfde kop 3 als in de markdown, uit `render_inleiding`
    assert _koppen(doc, "h3") == [sjabloon.BEDRIJFSTRAINING_KOP]


def test_docs_koppen_hebben_de_afgesproken_grootte_en_gewicht():
    """Kop 1 en 2 niet vet, kop 3 wel; 20/16/14pt. Zo staat het in het reviewdocument."""
    doc = uit.render_docs_html(uit.document_to_content(_document(), {}), "Cursus XML")
    for niveau, (grootte, gewicht, _) in uit.DOCS_KOPPEN.items():
        stijl = f"font-size:{grootte};font-weight:{gewicht};"
        assert f'<{niveau} style="{stijl}' in doc, (niveau, stijl)
        # ook op een <span> binnen de kop: Docs bewaart grootte en vet als tekenopmaak, en
        # laat het op de alinea alleen kan laten vallen
        assert f'<span style="{stijl}">' in doc, niveau
    assert uit.DOCS_KOPPEN["h1"][1] == "normal" and uit.DOCS_KOPPEN["h3"][1] == "bold"


def test_docs_html_geeft_elke_alinea_ruimte_eronder():
    """Zonder deze marge plakken alle alinea's van een kopje aan elkaar tot één blok."""
    doc = uit.render_docs_html(uit.document_to_content(_document(), {}), "Cursus XML")
    assert "<p>" not in doc, "een <p> zonder marge levert een doc zonder alinea's op"
    assert f'<p style="margin-top:0;margin-bottom:{uit.ALINEA_RUIMTE};">' in doc
    assert "<li>" not in doc and f"margin-bottom:{uit.BULLET_RUIMTE};" in doc
    # geen `margin`-shorthand en geen padding: zie `_ruimte`
    assert "margin:" not in doc and "padding" not in doc


def test_docopmaak_lekt_niet_naar_de_cms_content():
    """Het CMS levert zijn eigen opmaak; deze stijl hoort uitsluitend in het reviewdocument."""
    content = uit.document_to_content(_document(), {})
    for sleutel, waarde in content.items():
        if isinstance(waarde, str):
            assert "style=" not in waarde, f"{sleutel} draagt doc-opmaak het CMS in"
            assert "<span" not in waarde, sleutel
    # en de markdown evenmin
    assert "style=" not in uit.render_markdown(_document(), "Cursus XML")


def test_docs_html_zet_utf8_meta_in_de_kop():
    """Zonder deze meta gokt de importer latin-1 en gaat elke e-umlaut stuk."""
    doc = uit.render_docs_html(_content(), "Training Data")
    assert doc.startswith('<html><head><meta charset="utf-8">')
    assert doc.endswith("</body></html>")


def test_docs_html_zet_platte_kopjes_in_paragrafen():
    """`summary` en `summary_edudex` staan als platte tekst in het CMS (Kopje.html is False)."""
    content = _content(summary="Eerste zin.\n\nTweede zin.", summary_edudex="Kort.")
    doc = uit.render_docs_html(content, "Training Data")
    alineas = re.findall(r"<p\b[^>]*>(.*?)</p>", doc, re.S)
    assert "Eerste zin." in alineas and "Tweede zin." in alineas, alineas
    assert "Kort." in alineas


def test_docs_html_laat_geneste_bullets_uit_modules_intact():
    """De sub-bullets zijn de reden om HTML te uploaden en geen platte tekst."""
    doc = uit.render_docs_html(_content(), "Training Data")
    kaal = re.sub(r'\s*style="[^"]*"', "", doc)
    assert "<li>Module een<ul><li>Punt a</li>" in kaal


def test_docs_html_rendert_ook_een_overgenomen_training():
    """`neem_over` levert geen document, alleen content -- die trainingen moeten mee."""
    b = _briefing(modus_reviewer="overnemen", huidige_content=_content())
    res, content = rw.neem_over(b)
    assert not res.document, "vertrekpunt van deze test: het overnemen-spoor heeft geen document"
    doc = uit.render_docs_html(content, res.titel)
    assert _koppen(doc, "h2") == [k.kop for k in sjabloon.KOPJES]


def test_docs_html_geeft_een_leeg_kopje_toch_zijn_kop():
    """Een reviewer moet kunnen zien dat er niets staat, niet dat het kopje ontbreekt."""
    doc = uit.render_docs_html(_content(follow_up=""), "Training Data")
    assert "Vervolgstappen" in _koppen(doc, "h2")
    assert re.search(r"Vervolgstappen</span></h2><hr\b", doc), "het lege kopje kreeg toch inhoud"


# ---------------------------------------------------------------------------
# Artefacten op schijf: <id>.json + <id>.md
# ---------------------------------------------------------------------------

def _resultaat(document=None) -> rw.RewriteResult:
    return rw.RewriteResult(5, "Training XML", rw.APPROVED, document=document)


def test_artefacten_schrijven_json_en_markdown_naast_elkaar():
    doc = _document()
    with tempfile.TemporaryDirectory() as d:
        paden = rw.schrijf_training_artefacten(d, 5, _resultaat(doc), {"days": 3})
        assert sorted(os.listdir(d)) == ["5.json", "5.md"]
        with open(paden["md"], encoding="utf-8") as f:
            md = f.read()
        # identiek aan wat het notebook onder de cel toont
        assert md == uit.render_markdown(doc, "Training XML")
        with open(paden["json"], encoding="utf-8") as f:
            assert json.load(f)["content"] == {"days": 3}


def test_zonder_document_geen_markdown_en_een_oude_md_gaat_weg():
    """Een .md van een vorige run mag niet bij een nieuwere JSON blijven liggen."""
    with tempfile.TemporaryDirectory() as d:
        rw.schrijf_training_artefacten(d, 5, _resultaat(_document()), {})
        res = rw.RewriteResult(5, "Training XML", "error", reden="schrijver faalde")
        paden = rw.schrijf_training_artefacten(d, 5, res, {})
        assert paden["md"] is None
        assert os.listdir(d) == ["5.json"]


def test_artefact_dir_schrijft_plat_zonder_batch_en_in_een_submap_met():
    """Zonder batch blijft het de platte map: de trainingen van voor de indeling verhuizen niet."""
    assert rw.artefact_dir("uit") == os.path.join("uit", "trainingen")
    assert rw.artefact_dir("uit", "") == os.path.join("uit", "trainingen")
    assert rw.artefact_dir("uit", "  ") == os.path.join("uit", "trainingen")
    assert rw.artefact_dir("uit", "ronde 3") == os.path.join("uit", "trainingen", "ronde 3")


def test_zoek_artefact_vindt_een_training_in_elke_submap():
    """Een aanroeper weet niet in welke batch een training zit, en hoeft dat ook niet te weten."""
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 99)                      # plat
        _artefact(d, 2808, batch="ronde 3")   # in een submap
        assert rw.zoek_artefact(d, 99) == os.path.join(d, "trainingen", "99.json")
        assert rw.zoek_artefact(d, 2808) == os.path.join(d, "trainingen", "ronde 3", "2808.json")
        assert rw.zoek_artefact(d, 1234) is None


def test_artefact_paden_zoekt_zonder_batch_juist_wel_recursief():
    """`promoveer_naar_goud` kiest de few-shot uit alles wat we ooit hebben geschreven."""
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 99)
        _artefact(d, 1, batch="ronde 1")
        _artefact(d, 2, batch="ronde 2")
        alles = [os.path.basename(p) for p in rw.artefact_paden(d)]
        assert sorted(alles) == ["1.json", "2.json", "99.json"]
        assert [os.path.basename(p) for p in rw.artefact_paden(d, "ronde 1")] == ["1.json"]


def test_artefact_bewaart_de_flags_ook_uitgesplitst_per_tier():
    """Zonder `flags_tier` op schijf kan de Drive-comment de ruis niet van het oordeel scheiden."""
    res = rw.RewriteResult(5, "Training XML", rw.APPROVED, document=_document(),
                           flags=["hoog: fout", "laag: lang"],
                           flags_tier={checks.TIER_HOOG: ["hoog: fout"],
                                       checks.TIER_LAAG: ["laag: lang"]})
    with tempfile.TemporaryDirectory() as d:
        paden = rw.schrijf_training_artefacten(d, 5, res, {"days": 2})
        with open(paden["json"], encoding="utf-8") as f:
            op_schijf = json.load(f)
    assert op_schijf["flags_tier"][checks.TIER_HOOG] == ["hoog: fout"]
    assert op_schijf["flags"] == ["hoog: fout", "laag: lang"]


def test_overgenomen_training_landt_als_artefact_op_schijf():
    """Zonder dit bestaat een overgenomen training nergens op schijf, en mist hij op Drive.

    `neem_over` levert geen document, dus er komt geen .md -- maar wel de JSON met `content`,
    en dat is waar zowel sectie 7 als de doc-renderer op draait.
    """
    b = _briefing(modus_reviewer="overnemen", huidige_content=_content())
    res, content = rw.neem_over(b)
    with tempfile.TemporaryDirectory() as d:
        paden = rw.schrijf_training_artefacten(os.path.join(d, "trainingen"), 1, res, content)
        assert paden["md"] is None
        gevonden = drive.verzamel_uit_map(d)
    assert [t["training_id"] for t in gevonden] == [1]
    assert gevonden[0]["status"] == rw.OVERGENOMEN


def test_json_default_zet_numpy_scalars_om():
    import numpy as np
    assert rw._json_default(np.int64(2347)) == 2347
    assert isinstance(rw._json_default(np.int64(2347)), int)
    assert rw._json_default(np.bool_(True)) is True


def test_json_default_laat_echt_onserialiseerbare_objecten_falen():
    try:
        rw._json_default(object())
    except TypeError as e:
        assert "not JSON serializable" in str(e)
        return
    raise AssertionError("verwachtte TypeError voor een gewoon object")


def test_artefacten_verdragen_een_numpy_training_id():
    """Regressie: training_id komt uit een DataFrame en is dus numpy.int64, geen int."""
    import numpy as np
    res = rw.RewriteResult(np.int64(2347), "Training XML", rw.APPROVED, document=_document())
    res.thin = np.bool_(True)
    with tempfile.TemporaryDirectory() as d:
        paden = rw.schrijf_training_artefacten(d, res.training_id, res, {"days": np.int64(3)})
        assert sorted(os.listdir(d)) == ["2347.json", "2347.md"]
        with open(paden["json"], encoding="utf-8") as f:
            opgeslagen = json.load(f)
        assert opgeslagen["training_id"] == 2347
        assert type(opgeslagen["training_id"]) is int
        assert opgeslagen["thin"] is True
        assert opgeslagen["content"]["days"] == 3


def test_mislukt_schrijven_laat_de_vorige_versie_staan():
    """Een fout halverwege mag geen half artefact achterlaten (dat was de crash van cel 5)."""
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "5.json")
        with open(pad, "w", encoding="utf-8") as f:
            f.write('{"heel": "bestand"}')

        def stuk(f):
            f.write('{"half":')
            raise ValueError("simuleert een serialisatiefout")

        try:
            rw._schrijf_atomisch(pad, stuk)
        except ValueError:
            pass
        else:
            raise AssertionError("verwachtte dat de fout doorgegeven werd")

        with open(pad, encoding="utf-8") as f:
            assert json.load(f) == {"heel": "bestand"}
        assert os.listdir(d) == ["5.json"], "tmp-bestand niet opgeruimd"


def test_bewaar_training_zet_de_artefacten_in_trainingen():
    with tempfile.TemporaryDirectory() as d:
        paden = rw.bewaar_training(d, _resultaat(_document()), {"days": 7})
        assert paden["json"] == os.path.join(d, "trainingen", "5.json")
        with open(paden["json"], encoding="utf-8") as f:
            # dezelfde CMS-content als de batch schrijft, inclusief `days` uit de bron
            assert json.load(f)["content"]["days"] == 7


# ---------------------------------------------------------------------------
# De mate van aanpassing: twee assen (herschrijfniveau + actualiseringen)
# ---------------------------------------------------------------------------

def _content(**overrides) -> dict:
    """Bestaande CMS-content die aan het format voldoet, als vertrekpunt voor de scan."""
    basis = {
        "days": 2,
        "summary": "Wil je " + vul(57) + "?",
        "intro": "<p>" + vul(195, "onderwerp", "thema") + "</p>",
        "modules": ("<p>opening</p><ul>"
                    "<li>Module een<ul><li>Punt a</li><li>Punt b</li><li>Punt c</li></ul></li>"
                    "<li>Module twee<ul><li>Punt a</li><li>Punt b</li><li>Punt c</li>"
                    "<li>Punt d</li></ul></li>"
                    "<li>Module drie<ul><li>Punt a</li><li>Punt b</li><li>Punt c</li></ul></li>"
                    "<li>Module vier<ul><li>Punt a</li><li>Punt b</li><li>Punt c</li>"
                    "<li>Punt d</li><li>Punt e</li></ul></li>"
                    "</ul>"),
        "target_audience": "<p>Deze training is bedoeld voor iedereen die met data keuzes wil maken.</p>",
        "prior_knowledge": "<p>Specifieke voorkennis is niet noodzakelijk.</p>",
        "objectives": ("<p>Na deze training ben je in staat om:</p><ul>"
                       "<li>Datasets te ordenen en te controleren</li>"
                       "<li>Analyses te vertalen naar keuzes</li>"
                       "<li>Resultaten te presenteren aan je team</li>"
                       "<li>Vraagstukken gestructureerd te benaderen</li></ul>"),
        "summary_edudex": "Wil je slimmer met data werken en betere keuzes maken?",
        "setup": "<p>De training is praktijkgericht.</p>",
        "follow_up": "<p>Vervolg</p>",
        "certification": "<p>Certificaat</p>",
    }
    return {**basis, **overrides}


def test_modus_ladder_loopt_van_licht_naar_zwaar():
    assert rw.MODI == ("overnemen", "stijl", "format", "volledig")
    assert rw.MODUS_RANG["overnemen"] < rw.MODUS_RANG["stijl"] < rw.MODUS_RANG["format"]
    assert rw.hoogste_modus("stijl", "format") == "format"
    assert rw.hoogste_modus("volledig", "overnemen") == "volledig"
    # onbekende waarden vallen veilig terug, niet naar het lichtste niveau
    assert rw.normaliseer_modus("onzin") == rw.MODUS_DEFAULT == "volledig"
    assert rw.normaliseer_modus(float("nan")) == "volledig"
    assert rw.normaliseer_modus("  STIJL ") == "stijl"


def test_reviewer_modus_wint_van_het_voorstel():
    """Gezag volgt herkomst, net als bij de kern."""
    assert _briefing(modus_voorstel="format").modus == "format"
    assert _briefing(modus_voorstel="format", modus_reviewer="stijl").modus == "stijl"
    assert not _briefing(modus_voorstel="format").modus_van_reviewer
    assert _briefing(modus_reviewer="stijl").modus_van_reviewer


def test_herschreven_kolom_blijft_werken_zonder_de_nieuwe_kolommen():
    """Een scoresheet van vóór deze schaal moet zich exact gedragen als voorheen.

    `herschreven=1` betekende de facto al een modus ("niet aanraken"); zonder die terugval
    zou elke bestaande sheet ineens alles opnieuw laten schrijven.
    """
    assert _briefing(herschreven=True).modus == "overnemen"
    assert _briefing(herschreven=False).modus == "volledig"
    # een expliciet besluit gaat er wél overheen
    assert _briefing(herschreven=True, modus_reviewer="format").modus == "format"


def test_actualisering_verschuift_de_modus_niet():
    """As 2 staat los van as 1 -- dat is de kern van het ontwerp.

    Een goedgekeurde actie is een lokale toevoeging. Zou hij de modus ophogen, dan wordt de
    wijziging groter dan de reviewer vroeg: één toegevoegd onderwerp zou de hele training
    opnieuw laten schrijven.
    """
    actie = bes.Besluit(1, "T", 1, "Voeg GA4 toe", "ja", "doen", "", "handmatig")
    assert _briefing(modus_reviewer="stijl", goedgekeurd=[actie]).modus == "stijl"
    assert _briefing(herschreven=True, goedgekeurd=[actie]).modus == "overnemen"


def test_modules_parser_leest_beide_nestingsvormen():
    """De CMS-vorm zet de sub-lijst náást de titel-<li>, de renderer zet hem erin."""
    naast = ("<ul><li>Module een</li><ul><li>Punt a</li><li>Punt b</li></ul>"
             "<li>Module twee</li><ul><li>Punt c</li></ul></ul>")
    erin = ("<ul><li>Module een<ul><li>Punt a</li><li>Punt b</li></ul></li>"
            "<li>Module twee<ul><li>Punt c</li></ul></li></ul>")
    for html in (naast, erin):
        mods = rw._modules_uit_ul(html)
        assert [m["titel"] for m in mods] == ["Module een", "Module twee"], html
        assert [len(m["bullets"]) for m in mods] == [2, 1], html


def test_modules_in_h3_vorm_gelden_niet_als_leesbaar():
    """Onbekend is niet hetzelfde als conform: bij twijfel schat de scan omhoog."""
    h3 = ("<p>opening</p><h3>Module een</h3><p>Uitleg over de module.</p>"
          "<h3>Module twee</h3><p>Nog wat uitleg.</p>")
    assert not rw.modules_leesbaar(rw._modules_uit_ul(h3))
    assert rw.scan_vorm(_content(modules=h3), "Training X")["ondergrens"] == "format"


def test_scan_stelt_nooit_overnemen_voor():
    """De checks kunnen non-conformiteit bewijzen, conformiteit niet.

    Een tekst die elke regex haalt kan nog steeds het stijlregister negeren. Die conclusie
    mag daarom alleen van `schat_modus` of van een mens komen, nooit van de scan.
    """
    schoon = rw.scan_vorm(_content(), "Training Data")
    assert schoon["harde_issues"] == [], schoon["harde_issues"]
    assert schoon["ondergrens"] == "stijl"
    assert "niet door code vast te stellen" in schoon["reden"]


def test_scan_scheidt_structuur_van_formulering():
    """Ontbrekende kopjes en verkeerde aantallen -> format; verkeerde zinnen -> stijl."""
    # alleen de formulering deugt niet: de openingszin ontbreekt
    stijl = rw.scan_vorm(_content(summary="Deze training gaat over data."), "Training Data")
    assert stijl["ondergrens"] == "stijl"

    # een leeg verplicht kopje is structuur
    leeg = rw.scan_vorm(_content(objectives=""), "Training Data")
    assert leeg["ondergrens"] == "format"
    assert "doelen" in leeg["lege_kopjes"]

    # een verkeerd aantal sub-bullets ook
    krap = _content(modules="<ul><li>M1<ul><li>a</li></ul></li><li>M2<ul><li>b</li></ul></li></ul>")
    assert rw.scan_vorm(krap, "Training Data")["ondergrens"] == "format"


def test_scan_zonder_bron_of_met_onbruikbaar_verdict_gaat_naar_volledig():
    assert rw.scan_vorm({}, "T")["ondergrens"] == "volledig"
    assert rw.scan_vorm(_content(), "T", verdict="onbruikbaar")["ondergrens"] == "volledig"


def test_modus_voorstellen_stopt_op_een_id_zonder_bronrij():
    """Een id dat niet joint is een fout in het scoresheet, geen `volledig`-oordeel.

    Zou de lus doorlopen, dan zag `scan_vorm` lege content en kwam er het duurste advies uit
    dat er is -- op basis van een typefout. De poort hoort dus vóór de eerste call te zitten.
    """
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        bron_pad, scored_pad = os.path.join(d, "bron.xlsx"), os.path.join(d, "prio.xlsx")
        pd.DataFrame([{"id": 42, "name": "Training XML",
                       "content": json.dumps(_content(), ensure_ascii=False)}]
                     ).to_excel(bron_pad, index=False)
        _scored_df(id=[42, 999], name=["Training XML", "Training Zoek"],
                   actualiteit_actie=["", ""], actie_besluit=["", ""]
                   ).to_excel(scored_pad, index=False)
        try:
            rw.modus_voorstellen(scored_pad, bron_pad, met_llm=False, verbose=False)
        except ValueError as e:
            assert "999" in str(e) and "bronlijst" in str(e), str(e)
            return
    raise AssertionError("verwachtte ValueError bij een training_id zonder bronrij")


def test_modus_voorstellen_draait_door_als_elk_id_joint():
    """De poort mag een gezond sheet niet in de weg zitten -- ook zonder API-key."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        bron_pad, scored_pad = os.path.join(d, "bron.xlsx"), os.path.join(d, "prio.xlsx")
        pd.DataFrame([{"id": 42, "name": "Training XML",
                       "content": json.dumps(_content(), ensure_ascii=False)}]
                     ).to_excel(bron_pad, index=False)
        _scored_df(id=[42], name=["Training XML"],
                   actualiteit_actie=[""], actie_besluit=[""]).to_excel(scored_pad, index=False)
        uitkomst = rw.modus_voorstellen(scored_pad, bron_pad, met_llm=False, verbose=False)
        assert uitkomst["modus_voorstel"].iloc[0] in rw.MODI
        assert uitkomst["modules_nb_voorstel"].iloc[0] == sjabloon.MODULES_NB_DEFAULT
        assert "content is leeg" not in str(uitkomst["modus_reden"].iloc[0])


def test_schat_modus_valt_terug_op_de_ondergrens_zonder_client():
    """Levert het model niets, dan blijft staan wat de checks al hebben weerlegd."""
    uitkomst = rw.schat_modus(None, _content(objectives=""), "Training Data")
    assert uitkomst["modus"] == uitkomst["ondergrens"] == "format"


def test_schat_modus_mag_niet_onder_de_ondergrens_zakken():
    """Een te licht voorstel wordt opgehoogd, met de reden erbij."""
    def _client(modus):
        """Minimale stub die één submit_modus-tooluse teruggeeft."""
        blok = SimpleNamespace(type="tool_use", name="submit_modus",
                               input={"modus": modus, "reden": "ziet er goed uit"})
        resp = SimpleNamespace(content=[blok], stop_reason="tool_use")
        return SimpleNamespace(
            messages=SimpleNamespace(stream=lambda **_: _StubStroom(resp)))

    leeg = _content(objectives="")          # ondergrens = format
    op = rw.schat_modus(_client("overnemen"), leeg, "Training Data")
    assert op["modus"] == "format", op
    assert "opgehoogd naar de ondergrens" in op["reden"]
    # een zwaarder voorstel mag wél blijven staan
    zwaar = rw.schat_modus(_client("volledig"), leeg, "Training Data")
    assert zwaar["modus"] == "volledig"


def test_behoudende_modus_vervangt_de_brontekst_door_de_huidige_versie():
    """In stijl/format is de bestaande tekst het uitgangspunt, niet de brontekst.

    Twee keer dezelfde content meesturen kost tokens en geeft het model twee versies van de
    waarheid -- waarvan `source_text` de minst complete is (`setup`, `follow_up`,
    `summary_edudex` en `certification` ontbreken daar).
    """
    b = _briefing(modus_reviewer="stijl", huidige_content=_content())
    tekst = rw.build_writer_user(b)
    assert "HUIDIGE VERSIE" in tekst
    assert "BRONTEKST —" not in tekst
    assert "We introduceren het CRISP-DM model." not in tekst   # de source_text
    assert "redacteur, geen auteur" in tekst
    # de kopjes die build_source_text overslaat staan er nu wél in
    assert "Kortste omschrijving" in tekst

    # volledig blijft ongewijzigd: brontekst, geen huidige versie
    vol = rw.build_writer_user(_briefing(modus_reviewer="volledig", huidige_content=_content()))
    assert "BRONTEKST" in vol and "HUIDIGE VERSIE" not in vol
    assert vol.rstrip().endswith("We introduceren het CRISP-DM model.")


def test_zonder_bestaande_content_blijft_de_brontekst_staan():
    """Niets te behouden -> geen behoud-opdracht, ongeacht de modus."""
    tekst = rw.build_writer_user(_briefing(modus_reviewer="stijl", huidige_content={}))
    assert "BRONTEKST" in tekst and "HUIDIGE VERSIE" not in tekst


def test_actualiseer_alinea_verschijnt_alleen_met_goedgekeurde_acties():
    """Zonder deze alinea leest het model 'verander niets' als een verbod op de actie."""
    actie = bes.Besluit(1, "T", 1, "Voeg GA4 toe", "ja", "doen", "", "handmatig")
    met = rw.build_writer_user(_briefing(modus_reviewer="stijl", huidige_content=_content(),
                                         goedgekeurd=[actie]))
    assert "staan LOS van deze opdracht" in met
    zonder = rw.build_writer_user(_briefing(modus_reviewer="stijl", huidige_content=_content()))
    assert "staan LOS van deze opdracht" not in zonder


def test_judge_weet_in_welke_modus_hij_oordeelt():
    """Anders keurt hij een bewust conservatieve tekst af als 'te weinig herschreven'."""
    doc = {"titel": "T", "overzicht": "Wil je iets?"}
    stijl = rw.build_judge_user(_briefing(modus_reviewer="stijl", huidige_content=_content()), doc)
    assert "bewust ALLEEN bijgewerkt" in stijl
    assert "drift is hier een fail" in stijl
    # de bestaande instructies blijven staan; er komt alleen een as bij
    assert "Reken het concept NIET af op vorm" in stijl
    assert "FEITGETROUWHEID" in stijl and "NIVEAU" in stijl

    vol = rw.build_judge_user(_briefing(modus_reviewer="volledig"), doc)
    assert "bewust ALLEEN bijgewerkt" not in vol


def test_guidance_van_de_reviewer_gaat_achter_die_van_de_scorer():
    """De legacy-kolom `guidance_reviewer` blijft gelezen, gelabeld en achteraan.

    Sheets van voor augustus 2026 hebben de aanwijzing in die tweede kolom staan; sindsdien
    stelt de reviewer `rewrite_guidance` zelf bij. Beide moeten bij de schrijver aankomen.
    """
    b = _briefing(rewrite_guidance="Leg de nadruk op governance.",
                  guidance_reviewer="Modules 3 en 4 samenvoegen.")
    assert b.guidance_definitief.index("governance") < b.guidance_definitief.index("samenvoegen")
    tekst = rw.build_writer_user(b)
    assert "Modules 3 en 4 samenvoegen." in tekst
    assert "AANWIJZING VAN DE REVIEWER" in tekst
    # zonder reviewer-aanwijzing verandert er niets aan de bestaande regel
    assert "AANWIJZING VAN DE REVIEWER" not in rw.build_writer_user(_briefing())


def test_bijgestelde_guidance_bereikt_ook_de_actualisering():
    """Een aanwijzing die de reviewer in `rewrite_guidance` typt, komt ook op het overnemen-pad.

    Dit pad las alleen `guidance_reviewer`, en dat kon toen dat een eigen kolom was. Nu de vrije
    aanwijzing in `rewrite_guidance` staat, zou die lezing elke aanwijzing van het reviewteam
    laten vallen -- precies bij de trainingen waar de reviewer de enige is die iets vraagt.
    """
    b = _briefing(rewrite_guidance="Noem de nieuwe wetgeving bij naam.",
                  goedgekeurd=[bes.Besluit(1, "T", 1, "refresh: wetgeving", "prima",
                                           bes.DOEN, "", bes.BRON_REGEL, "")])
    tekst = rw.build_actualisatie_user(b, _content(), "Training X")
    assert "Noem de nieuwe wetgeving bij naam." in tekst
    # ... maar met de grens eromheen: `overnemen` mag geen herstructurering worden
    assert "GEEN opdracht om de" in tekst
    assert "AANWIJZING bij deze training" not in rw.build_actualisatie_user(
        _briefing(goedgekeurd=b.goedgekeurd), _content(), "Training X")


def test_actualisatie_tool_heeft_geen_verplichte_kopjes():
    """Het model levert alleen wat verandert; de rest blijft byte-voor-byte staan."""
    tool = rw.build_actualisatie_tool()
    assert tool["input_schema"]["required"] == []
    assert set(rw.ACTUALISEERBARE_KOPJES) <= set(tool["input_schema"]["properties"])
    # de veldbeschrijvingen komen uit SUBMIT_REWRITE, zodat ze niet uiteen lopen
    assert (tool["input_schema"]["properties"]["overzicht"]
            is rw.SUBMIT_REWRITE["input_schema"]["properties"]["overzicht"])


def test_render_veld_raakt_alleen_het_eigen_cms_veld():
    sleutel, waarde = uit.render_veld("doelgroep", "Deze training is bedoeld voor analisten.")
    assert sleutel == "target_audience"
    assert waarde == "<p>Deze training is bedoeld voor analisten.</p>"
    sleutel, waarde = uit.render_veld("overzicht", "Wil je iets?")
    assert sleutel == "summary" and waarde == "Wil je iets?"      # platte tekst, geen HTML


def test_overnemen_zonder_acties_doet_geen_enkele_call():
    """Het goedkope pad moet gratis blijven als er niets te actualiseren valt."""
    b = _briefing(titel="Cursus SQL", huidige_content=_content(), herschreven=True)
    res, content = rw.neem_over(b, client=None)
    assert res.status == rw.OVERGENOMEN
    assert res.modus == "overnemen"
    assert res.titel == "Training SQL"           # alleen de titelnormalisatie
    assert content["summary"] == _content()["summary"]
    assert res.toegepaste_acties == []


def test_overnemen_meldt_het_wanneer_een_actie_niet_kon_worden_doorgevoerd():
    """Stilzwijgend laten vallen is precies de fout die deze schaal moest repareren."""
    actie = bes.Besluit(1, "T", 1, "Voeg GA4 toe", "ja", "doen", "", "handmatig")
    b = _briefing(titel="Training SQL", huidige_content=_content(), herschreven=True,
                  goedgekeurd=[actie])
    res, _ = rw.neem_over(b, client=None)     # geen client -> niets doorgevoerd
    assert any("geen wijziging" in f for f in res.flags), res.flags
    assert res.toegepaste_acties == ["1. Voeg GA4 toe"]


def test_resultaat_legt_vast_onder_welk_regime_het_tot_stand_kwam():
    """Zonder modus en spec-versie is `approved` een status zonder betekenis."""
    b = _briefing(huidige_content=_content(), herschreven=True)
    res, content = rw.neem_over(b, client=None)
    rij = rw._review_rij(res, content)
    assert rij["modus"] == "overnemen"
    assert "modus_voorstel" in rij and "spec_versie" in rij
    assert len(rw.spec_versie()) == 12
    # de vingerafdruk verandert mee met de spec, niet met de code
    assert rw.spec_versie() == rw.spec_versie()


# ---------------------------------------------------------------------------
# Templatev2: vaste teksten, Modules-varianten, lopende aanduiding
# ---------------------------------------------------------------------------

def test_schrijfspec_citeert_de_actuele_vaste_teksten():
    """De schrijfspec citeert de vaste teksten letterlijk; die twee mogen niet uit elkaar lopen.

    Dit is de enige echte bewaker tegen de drievoudige duplicatie (sjabloon.py, de schrijfspec
    en het template). Verandert er een vaste tekst en vergeet iemand de spec, dan leest de
    schrijver een andere zin dan de code invoegt -- en dat is precies het soort verschil dat
    pas in de output opvalt.
    """
    spec = open(rw.SCHRIJFSPEC, encoding="utf-8").read()
    # De spec citeert de vaste teksten in een blockquote, en een tekst met een alineagrens erin
    # (de Modules-openingszin, met de NB als tweede alinea) krijgt daardoor een ">" middenin.
    # Dat is markdown-opmaak en geen inhoudsverschil, dus die haalt deze test eraf.
    spec = re.sub(r"(?m)^\s*>\s?", "", spec)
    spec = " ".join(spec.split())
    # Alleen de teksten die de spec daadwerkelijk citeert; de placeholder-varianten toetsen
    # we op hun deel vóór de {}.
    verplicht = [
        sjabloon.BEDRIJFSTRAINING_KOP,
        # inclusief de aanhalingstekens: de spec hoort de nadruk te tonen die de code plaatst
        sjabloon.AANPAK_ALINEA_2,
        sjabloon.VERVOLG_ALINEA_1,
        sjabloon.VERVOLG_ALINEA_2,
        sjabloon.VERVOLG_LIJST_INTRO,
        sjabloon.DOELEN_INTRO,
        sjabloon.VOORKENNIS_FALLBACK,
        sjabloon.MODULES_NB_STABIEL.split("{aanduiding}")[1].strip(),
        sjabloon.MODULES_NB_ACTUEEL.split("{aanduiding}")[1].strip(),
        sjabloon.AANPAK_ALINEA_1.split("{invulling}")[0].strip(),
    ]
    ontbreekt = [t for t in verplicht if " ".join(t.split()) not in spec]
    assert not ontbreekt, f"schrijfspec citeert verouderde vaste tekst: {ontbreekt}"


def test_lopende_aanduiding_zet_soortwoord_in_onderkast():
    assert sjabloon.lopende_aanduiding("Training PHP Professional") == "de training PHP Professional"
    assert sjabloon.lopende_aanduiding("Cursus XML") == "de training XML"
    assert sjabloon.lopende_aanduiding("Masterclass PHP") == "de masterclass PHP"
    assert sjabloon.lopende_aanduiding("Examentraining CEH") == "de examentraining CEH"
    assert sjabloon.lopende_aanduiding("Power BI") == "de training Power BI"
    assert sjabloon.lopende_aanduiding("") == "deze training"
    # de rest van de titel houdt zijn hoofdletters
    assert sjabloon.lopende_aanduiding("Training Power BI") == "de training Power BI"


def test_modules_nb_varianten():
    stabiel = sjabloon.modules_opening("Training XML")
    actueel = sjabloon.modules_opening("Training XML", "actueel")
    assert "Mocht je vragen hebben" in stabiel
    assert "snelle ontwikkelingen" in actueel
    # default is de terughoudende variant
    assert sjabloon.modules_opening("Training XML", "") == stabiel
    assert sjabloon.modules_opening("Training XML", "onzin") == stabiel
    # "in basis" is uit beide verdwenen
    assert "in basis" not in stabiel and "in basis" not in actueel


def test_normaliseer_modules_nb():
    assert rw.normaliseer_modules_nb("actueel") == "actueel"
    assert rw.normaliseer_modules_nb("ACTUEEL") == "actueel"
    assert rw.normaliseer_modules_nb("") == "stabiel"
    assert rw.normaliseer_modules_nb(None) == "stabiel"
    assert rw.normaliseer_modules_nb("misschien") == "stabiel"


def test_briefing_modules_nb_gezagsvolgorde():
    b = rw.RewriteBriefing(training_id=1, titel="Training X", persona="A", dagen=2,
                           kern="k", verdict="rijk", actualiteit_type="", source_text="s")
    assert b.modules_nb == "stabiel"                       # niets ingevuld
    b.modules_nb_voorstel = "actueel"
    assert b.modules_nb == "actueel"                       # voorstel telt
    b.modules_nb_reviewer = "stabiel"
    assert b.modules_nb == "stabiel"                       # reviewer wint van voorstel


def test_afsluiter_vervalt_zonder_lege_alinea():
    assert sjabloon.VERVOLG_AFSLUITER == ""
    html_out = uit.render_vervolgstappen(
        [sjabloon.VERVOLG_ALINEA_1], ["Power BI"], sjabloon.VERVOLG_AFSLUITER)
    assert "<p></p>" not in html_out
    assert not html_out.rstrip().endswith("<p></p>")


def test_deelnamecertificaat_kop_met_ongewijzigd_cms_veld():
    kopje = sjabloon.KOPJES[-1]
    assert kopje.kop == "Deelnamecertificaat"
    # het CMS-contract mag NIET meeveranderen met de kopnaam
    assert kopje.veld == "certificatie" and kopje.cms == "certification"


# ---------------------------------------------------------------------------
# Nieuwe checks: doelgroep-opening, lerend aspect, soortwoord-hoofdletter
# ---------------------------------------------------------------------------

def test_doelgroep_eist_bedoeld_voor():
    rwd = _good_rewrite()
    rwd["doelgroep"] = "Deze training is voor iedereen die met data werkt."
    codes = [i.code for i in hard_fails(check_rewrite(rwd))]
    assert "opening" in codes
    rwd["doelgroep"] = "Deze training is bedoeld voor iedereen die met data werkt."
    assert "opening" not in [i.code for i in hard_fails(check_rewrite(rwd))]


def test_lerend_aspect_is_flag_geen_hardfail():
    # het schone concept demonstreert het lerende aspect en vuurt dus niet
    assert "lerend_aspect" not in _codes(check_rewrite(_good_rewrite(), _CTX))
    rwd = _good_rewrite()
    rwd["overzicht"] = "Wil je datamodellen opzetten die kloppen? " + vul(55, "woord", "term")
    issues = check_rewrite(rwd)
    assert "lerend_aspect" in [i.code for i in flags(issues)]
    assert "lerend_aspect" not in [i.code for i in hard_fails(issues)]


def test_lerend_aspect_accepteert_nederlandse_vormen():
    import rewrite_checks as c
    for zin in ("Wil je leren modelleren?", "Wil je dit kunnen opzetten?",
                "Je leert data te beoordelen.", "Wil je in staat zijn om te modelleren?",
                "Wil je weten hoe je dit aanpakt?"):
        assert c._LEREND_RE.search(zin), zin
    assert not c._LEREND_RE.search("Wil je datamodellen opzetten?")


def test_soortwoord_hoofdletter_is_flag():
    rwd = _good_rewrite()
    rwd["inleiding"] = "Tijdens de Training XML leer je veel. " + vul(180, "woord", "term")
    issues = check_rewrite(rwd)
    assert "soortwoord_hoofdletter" in [i.code for i in flags(issues)]
    assert "soortwoord_hoofdletter" not in [i.code for i in hard_fails(issues)]


# ---------------------------------------------------------------------------
# Vaste teksten verversen (het `overnemen`-pad)
# ---------------------------------------------------------------------------

_OUDE_CONTENT = {
    "intro": "<p>Geschreven inleiding.</p>\n\n"
             "<h3>Deze training bieden we ook als bedrijfstraining voor jou en je team</h3>\n"
             "<p>De inhoud stemmen we dan af op jullie werksituatie, systemen en concrete "
             "vraagstukken, zodat de training direct aansluit.</p>",
    "modules": "<p>Tijdens de Training XML komen in basis onderstaande onderwerpen aan bod. "
               "Afhankelijk van ontwikkelingen op het vakgebied, kan de feitelijke "
               "trainingsinhoud hier echter van afwijken.</p>\n\n"
               "<ul>\n  <li>Module A\n    <ul><li>bullet</li></ul>\n  </li>\n</ul>",
    "setup": "<p>De training is interactief en praktijkgericht opgezet. Je werkt actief aan "
             "herkenbare situaties. Door te oefenen, bespreken en reflecteren ervaar je hoe "
             "je XML in de praktijk toepast.</p>\n\n"
             "<p>De training wordt verzorgd door trainers uit de praktijk.</p>",
    "follow_up": "<p>Binnen dit vakgebied beschikken wij over ruime praktijkervaring.</p>\n\n"
                 "<p>Er zijn verschillende vervolgtrainingen die aansluiten.</p>\n\n"
                 "<p>Wil je je verdiepen, dan sluiten deze aan:</p>\n<ul>\n<li>XSLT</li>\n</ul>"
                 "\n\n<p>Zo kies je een vervolgstap die past bij jouw rol, interesses en "
                 "werksituatie.</p>",
    "certification": "<p>Na het volledig afronden van deze training ontvang je een certificaat "
                     "van deelname.</p>",
}


def test_ververs_vaste_teksten_vervangt_alle_vijf():
    nieuw, gewijzigd = uit.ververs_vaste_teksten(_OUDE_CONTENT, "Training XML")
    assert [g for g in gewijzigd if not g.startswith("witregels")] == [
        "Deelnamecertificaat", "bedrijfstrainingblok", "Modules-openingszin", "Aanpak",
        "Vervolgstappen-boilerplate"], gewijzigd
    for veld in ("intro", "modules", "setup", "follow_up", "certification"):
        plat = uit._tekst_uit(nieuw[veld])
        for vervallen in sjabloon.VERVALLEN_VASTE_TEKSTEN:
            assert vervallen not in plat, f"{veld} houdt nog: {vervallen}"


def test_ververs_behoudt_geschreven_tekst():
    nieuw, _ = uit.ververs_vaste_teksten(_OUDE_CONTENT, "Training XML")
    # de geschreven inleiding, de modulelijst, de Aanpak-invulling en de catalogustitels
    assert "Geschreven inleiding." in nieuw["intro"]
    assert "Module A" in nieuw["modules"] and "<li>bullet</li>" in nieuw["modules"]
    assert "je XML in de praktijk toepast" in nieuw["setup"]
    assert "<li>XSLT</li>" in nieuw["follow_up"]
    assert "Wil je je verdiepen, dan sluiten deze aan:" in nieuw["follow_up"]


def test_ververs_is_idempotent():
    eenmaal, _ = uit.ververs_vaste_teksten(_OUDE_CONTENT, "Training XML")
    tweemaal, gewijzigd = uit.ververs_vaste_teksten(eenmaal, "Training XML")
    assert gewijzigd == [], gewijzigd
    assert eenmaal == tweemaal


def test_ververs_respecteert_modules_variant():
    nieuw, _ = uit.ververs_vaste_teksten(_OUDE_CONTENT, "Training XML", "actueel")
    assert "snelle ontwikkelingen" in nieuw["modules"]


def test_nb_staat_in_een_eigen_alinea():
    """Drie reviewers schreven onafhankelijk "nieuwe alinea" bij de NB onder Modules.

    Een tweede <p> is precies hoe de bestaande CMS-content zijn alinea's maakt (de `intro` van
    69 van de 78 trainingen heeft er drie of vier). Wat níét mag is een letterlijke newline
    ertussen -- dat is de bevinding uit de vorige ronde.
    """
    html = uit.render_modules(sjabloon.modules_opening("Training XML"),
                              [{"titel": "Module een", "bullets": ["Onderdeel a"]}])
    kop = html.split("<ul", 1)[0]
    assert kop.count("<p>") == 2, kop
    assert "</p><p>NB:" in kop
    assert "\n" not in html


def test_stabiele_nb_vraagt_naar_de_inhoud_niet_naar_de_actuele_inhoud():
    """De actualiteit is precies wat de variant `actueel` afdekt; in allebei roept het een
    vraag op die de stabiele variant zelf niet beantwoordt."""
    assert "over de inhoud" in sjabloon.MODULES_NB_STABIEL
    assert "actuele inhoud" not in sjabloon.MODULES_NB_STABIEL
    assert "actuele inhoud" in sjabloon.MODULES_NB_ACTUEEL


def test_ververs_zet_de_nb_alinea_om_en_laat_de_modulelijst_staan():
    """Op echte CMS-content: de opening wordt opnieuw opgebouwd, de lijst blijft ongemoeid."""
    bron = _goud_content()
    if not bron.get("modules"):
        return
    nieuw, _ = uit.ververs_vaste_teksten(bron, "Training XML")
    kop = nieuw["modules"].split("<ul", 1)[0]
    assert kop.count("<p>") == 2, kop
    assert "actuele inhoud" not in kop
    # alles vanaf de lijst is ongewijzigd, op de witruimtenormalisatie na
    oud_lijst = bron["modules"][bron["modules"].index("<ul"):]
    nieuw_lijst = nieuw["modules"][nieuw["modules"].index("<ul"):]
    assert re.sub(r"\s+", "", oud_lijst) == re.sub(r"\s+", "", nieuw_lijst)


def test_aanpak_alinea_2_noemt_het_expertisegebied():
    """Reviewronde 2: de boilerplate liep achter op de vaste woordenschat uit Sectie 0.20."""
    assert sjabloon.AANPAK_ALINEA_2.startswith(
        "Onze trainers zijn, naast trainer, dagelijks werkzaam op dit expertisegebied.")
    assert "expert op hun trainingsonderwerp" not in sjabloon.AANPAK_ALINEA_2


# ---------------------------------------------------------------------------
# Reviewbevindingen op de eerste batch (796 / 2347 / 2407)
# ---------------------------------------------------------------------------

def _goud_content(bestandsnaam: str = "107.json") -> dict:
    """Echte CMS-content van een training die al in de nieuwe stijl staat."""
    pad = os.path.join(rw.GOUD_DIR, bestandsnaam)
    if not os.path.exists(pad):
        return {}
    with open(pad, encoding="utf-8") as f:
        return json.load(f).get("content") or {}


def test_bestaande_cms_content_bevat_geen_newlines():
    """De referentie waar de output op gemeten wordt -- valt die weg, dan zegt de test hieronder niets."""
    content = _goud_content()
    if not content:
        return
    for sleutel, waarde in content.items():
        if isinstance(waarde, str):
            assert "\n" not in waarde, f"{sleutel} bevat wél een newline"


def test_output_html_bevat_geen_letterlijke_newlines():
    """Alinea's worden gescheiden door <p>-tags, niet door witregels in de HTML.

    Het CMS zet die letterlijke newlines om in extra witruimte, waardoor de tekst er anders
    uitziet dan de trainingen die er al in staan -- en die hebben er geen enkele.
    """
    content = uit.document_to_content(_document(), {})
    for sleutel, waarde in content.items():
        if isinstance(waarde, str):
            assert "\n" not in waarde, f"{sleutel}: {waarde[:120]!r}"


def test_ververs_haalt_oude_boilerplate_uit_echte_cms_content():
    """Het `overnemen`-pad op content zoals die écht in het CMS staat: zonder newlines.

    De eerdere split op een witregel leverde daar één blok op, waardoor de oude vaste alinea's
    bleven staan en de nieuwe eroverheen werden geplakt -- dubbele boilerplate dus.
    """
    content = _goud_content()
    if not content.get("follow_up"):
        return
    nieuw, gewijzigd = uit.ververs_vaste_teksten(content, "Training Creative Cloud")
    assert "Vervolgstappen-boilerplate" in gewijzigd
    plat = uit._tekst_uit(nieuw["follow_up"])
    assert "Wil je je na deze training verder verdiepen of verbreden?" not in plat
    assert "Binnen dit vakgebied beschikken wij" not in plat
    assert plat.count(sjabloon.VERVOLG_ALINEA_2) == 1
    # de catalogustitels en hun eigen intro's blijven staan
    assert "Training DTP met InDesign" in nieuw["follow_up"]


def test_witregelnormalisatie_plakt_geen_woorden_aan_elkaar():
    """Alleen witruimte rond de blokstructuur mag weg; tussen twee woorden nooit."""
    assert uit._compacte_html("<p>een zin</p>\n\n<p>en nog een</p>") == \
        "<p>een zin</p><p>en nog een</p>"
    assert uit._compacte_html("<ul>\n  <li>Titel\n    <ul>\n      <li>bullet</li>\n</ul>\n</li>\n</ul>") == \
        "<ul><li>Titel<ul><li>bullet</li></ul></li></ul>"
    # spaties binnen een tekstknoop blijven
    assert uit._compacte_html("<p>twee woorden <em>schuin</em> erna</p>") == \
        "<p>twee woorden <em>schuin</em> erna</p>"


def test_aanpak_invulling_verdubbelt_hoe_niet():
    """Training 2347 leverde "hoe een data-analysetraject ..." en kreeg "ervaar je hoe hoe"."""
    briefing = _briefing()
    doc = rw.assemble_document({**_good_rewrite(),
                                "aanpak_invulling": "hoe een data-analysetraject verloopt"},
                               briefing, [])
    assert "ervaar je hoe hoe" not in doc["aanpak"]
    assert "ervaar je hoe een data-analysetraject verloopt." in doc["aanpak"]


def test_aanpak_invulling_met_voegwoord_is_een_flag():
    rwd = _good_rewrite()
    rwd["aanpak_invulling"] = "hoe je datamodellen opzet"
    issues = check_rewrite(rwd, _CTX)
    assert "invulling_voegwoord" in [i.code for i in flags(issues)]
    assert "invulling_voegwoord" not in [i.code for i in hard_fails(issues)]


def test_hoe_hoe_overleeft_een_round_trip_niet():
    """De fout was zelfbestendigend: de terugleespaden vingen "hoe X" en plakten hem terug."""
    kapot = sjabloon.AANPAK_ALINEA_1.format(invulling="hoe een traject verloopt")
    hersteld = rw._writer_out_uit_json({"document": {"aanpak": kapot}})
    assert hersteld["aanpak_invulling"] == "een traject verloopt"
    opnieuw, _ = uit.ververs_vaste_teksten(
        {"setup": f"<p>{kapot}</p>"}, "Training XML")
    assert "ervaar je hoe hoe" not in opnieuw["setup"]


def test_dubbel_woord_is_hardfail_maar_je_je_mag():
    rwd = _good_rewrite()
    rwd["doelgroep"] = "Deze training is bedoeld voor iedereen die met met data werkt."
    assert "dubbel_woord" in [i.code for i in hard_fails(check_rewrite(rwd, _CTX))]
    # "maak je je de materie eigen" staat in onze eigen vaste tekst en is correct Nederlands
    rwd["doelgroep"] = "Deze training is bedoeld voor iedereen die zich je je werk eigen maakt."
    assert "dubbel_woord" not in [i.code for i in hard_fails(check_rewrite(rwd, _CTX))]


def test_voorkennis_mag_twee_zinnen_zijn():
    """Het aanbevolen antwoord uit schrijfspec Sectie 7 bestaat zelf uit twee zinnen."""
    rwd = _good_rewrite()
    rwd["voorkennis"] = ("Enige ervaring in het werken met JavaScript is vereist. Mocht je "
                         "hier vragen over hebben, neem gerust contact met ons op.")
    issues = check_rewrite(rwd, _CTX)
    assert "een_zin" not in _codes_in(issues, "voorkennis")
    assert "voorkennis_lang" not in _codes_in(issues, "voorkennis")


def test_voorkennis_die_uitloopt_is_wel_een_flag():
    rwd = _good_rewrite()
    rwd["voorkennis"] = "Enige ervaring met data is vereist. " + vul(50)
    assert "voorkennis_lang" in _codes_in(check_rewrite(rwd, _CTX), "voorkennis", FLAG)


def test_doelgroep_blijft_wel_op_een_zin_staan():
    """Daar is één zin wél het ontwerp; alleen Voorkennis is versoepeld."""
    rwd = _good_rewrite()
    rwd["doelgroep"] = "Deze training is bedoeld voor analisten. Ook voor adviseurs."
    assert "een_zin" in _codes_in(check_rewrite(rwd, _CTX), "doelgroep", FLAG)


# ---------------------------------------------------------------------------
# Reviewronde 2: anglicismen, formuleringen aan de onderkant, "Na deze training",
# de contactzin en de NB-alinea onder Modules
# ---------------------------------------------------------------------------

def test_anglicisme_is_flag_geen_hardfail():
    """Expliciet gevraagd door de schrijfstijl-eigenaar. Flag: het kan een vakterm zijn."""
    rwd = _good_rewrite()
    rwd["inleiding"] = "Daarna werk je door de categorieën heen. " + vul(190, "onderwerp", "thema")
    issues = check_rewrite(rwd, _CTX)
    assert "anglicisme" in _codes_in(issues, "inleiding", FLAG)
    assert "anglicisme" not in _codes(issues, HARD)


def test_leenwoord_met_nederlandse_tegenhanger_is_flag():
    rwd = _good_rewrite()
    rwd["doelgroep"] = "Deze training is bedoeld voor iedereen die zijn skills wil uitbouwen."
    issue = [i for i in check_rewrite(rwd, _CTX) if i.code == "anglicisme"]
    assert issue and "vaardigheden" in issue[0].message


def test_vaktermen_uit_de_eigen_catalogus_vuren_niet():
    """De lijst is op `herschreven/goud/` gekalibreerd: wat een echte vakterm kan zijn blijft
    eruit. "best practices" staat zelfs in onze eigen schrijfspec Sectie 1a."""
    rwd = _good_rewrite()
    rwd["voorkennis"] = ("Enige ervaring met governance, compliance, deployment, performance "
                         "en best practices is aan te raden.")
    assert "anglicisme" not in _codes_in(check_rewrite(rwd, _CTX), "voorkennis", FLAG)


def test_hooguit_een_anglicisme_per_veld():
    """Een tekst met drie leenwoorden heeft één probleem, niet drie."""
    rwd = _good_rewrite()
    rwd["doelgroep"] = "Deze training is bedoeld voor iedereen met skills, insights en een mindset."
    issues = [i for i in check_rewrite(rwd, _CTX) if i.code == "anglicisme"]
    assert len(issues) == 1


def test_zwakke_formulering_is_flag():
    """De grootste groep uit ronde 2: het niveau klopt, het werkwoord staat aan de onderkant."""
    for veld, tekst in (
        ("overzicht", "Wil je de begrippen rond data kunnen plaatsen? " + vul(52)),
        ("inleiding", "Je ervaart hoe een analysetraject in elkaar zit. "
                      + vul(188, "onderwerp", "thema")),
        ("doelgroep", "Deze training is bedoeld voor iedereen die gerichter wil meepraten."),
    ):
        rwd = _good_rewrite()
        rwd[veld] = tekst
        assert "zwakke_formulering" in _codes_in(check_rewrite(rwd, _CTX), veld, FLAG), veld


def test_sterke_formulering_binnen_de_scope_vuurt_niet():
    """Wat de schrijfstijl-eigenaar wél goedkeurde voor een foundation-training."""
    rwd = _good_rewrite()
    rwd["overzicht"] = ("Wil je de opbouw van een analysetraject kunnen doorgronden en een "
                        "stevige basis kunnen leggen als analist? " + vul(43))
    assert "zwakke_formulering" not in _codes_in(check_rewrite(rwd, _CTX), "overzicht", FLAG)


def test_tweede_zin_zonder_in_deze_training_is_flag():
    rwd = _good_rewrite()
    rwd["overzicht"] = "Wil je data kunnen " + vul(30) + "? Je leert " + vul(24) + "."
    issues = check_rewrite(rwd, _CTX)
    assert "tweede_zin" in _codes_in(issues, "overzicht", FLAG)
    assert "tweede_zin" not in _codes(issues, HARD)


def test_tweede_zin_met_in_deze_training_vuurt_niet():
    rwd = _good_rewrite()
    rwd["overzicht"] = ("Wil je data kunnen " + vul(30) + "? In deze training leer je "
                        + vul(20) + ".")
    assert "tweede_zin" not in _codes_in(check_rewrite(rwd, _CTX), "overzicht", FLAG)


def test_tweede_zin_accepteert_masterclass_en_een_bijvoeglijk_naamwoord():
    """§0.0 laat 'Masterclass'/'Workshop' staan, en 'In deze interactieve training' is prima."""
    for opening in ("Tijdens deze masterclass ", "In deze interactieve training ",
                    "Tijdens de workshop "):
        rwd = _good_rewrite()
        rwd["overzicht"] = "Wil je data kunnen " + vul(30) + "? " + opening + vul(22) + "."
        assert "tweede_zin" not in _codes_in(check_rewrite(rwd, _CTX), "overzicht", FLAG), opening


def test_overzicht_van_een_zin_levert_geen_tweede_zin_flag():
    """Geen tweede zin = niets om over te oordelen; de lengtecheck vangt dat al."""
    import rewrite_checks as c
    rwd = _good_rewrite()
    assert len(c.zinnen(rwd["overzicht"])) == 1
    assert "tweede_zin" not in _codes_in(check_rewrite(rwd, _CTX), "overzicht", FLAG)


def test_kortste_omschrijving_zonder_na_deze_training_is_flag():
    rwd = _good_rewrite()
    rwd["kortste_omschrijving"] = "Wil je slimmer met data kunnen werken? Je leert het stap voor stap."
    issues = check_rewrite(rwd, _CTX)
    assert "geen_na_deze_training" in _codes_in(issues, "kortste_omschrijving", FLAG)
    assert "geen_na_deze_training" not in _codes(issues, HARD)


def test_na_afloop_van_deze_training_telt_ook():
    rwd = _good_rewrite()
    rwd["kortste_omschrijving"] = ("Wil je slimmer met data kunnen werken? Na afloop van deze "
                                   "training onderbouw je je keuzes met cijfers.")
    assert "geen_na_deze_training" not in _codes_in(
        check_rewrite(rwd, _CTX), "kortste_omschrijving", FLAG)


def test_contactzin_zonder_dan_is_flag():
    rwd = _good_rewrite()
    rwd["voorkennis"] = ("Enige ervaring met data is vereist. Mocht je hier vragen over hebben, "
                         "neem gerust contact met ons op.")
    assert "contactzin_zonder_dan" in _codes_in(check_rewrite(rwd, _CTX), "voorkennis", FLAG)


def test_contactzin_met_dan_vuurt_niet():
    rwd = _good_rewrite()
    rwd["voorkennis"] = ("Enige ervaring met data is vereist. Mocht je hier vragen over hebben, "
                         "neem dan gerust contact met ons op.")
    assert "contactzin_zonder_dan" not in _codes_in(check_rewrite(rwd, _CTX), "voorkennis", FLAG)


def test_duur_in_de_tekst_is_hardfail():
    for veld, tekst in (
        ("inleiding", "In deze training van twee dagen leer je veel. " + vul(180, "onderwerp", "thema")),
        ("overzicht", "Wil je in een tweedaagse training data leren duiden? " + vul(52)),
        ("kortste_omschrijving", "Wil je in 2 dagen leren modelleren?"),
    ):
        rwd = _good_rewrite()
        rwd[veld] = tekst
        codes = _codes_in(check_rewrite(rwd, _CTX), veld, HARD)
        assert "duur_in_tekst" in codes, f"{veld}: {codes}"


def test_duur_check_raakt_geen_module_inhoud():
    """"Tweedaagse implementatie" in een bullet kan over de stof gaan, niet over ons."""
    rwd = _good_rewrite()
    rwd["modules"]["modules"][0]["bullets"][0] = "Een tweedaagse implementatie voorbereiden"
    assert "duur_in_tekst" not in [i.code for i in check_rewrite(rwd, _CTX)]


def test_modulesband_schuift_mee_met_de_duur():
    import rewrite_checks as c
    assert c.modulesband(1) == (4, 6)
    assert c.modulesband(2) == (4, 6)
    assert c.modulesband(3) == (4, 6)
    assert c.modulesband(5) == (5, 8)
    assert c.modulesband(None) == (4, 6)   # dagen onbekend -> de smalle middenband


def _n_modules(n: int) -> list[dict]:
    return [{"titel": f"Module {i}", "bullets": ["a", "b", "c"] + (["d"] if i % 2 else [])}
            for i in range(n)]


def test_zeven_modules_mag_bij_vijf_dagen_maar_niet_bij_drie():
    """De band is teruggeschroefd: zeven modules bij een tweedaagse is nu een harde fail."""
    rwd = _good_rewrite()
    rwd["modules"]["modules"] = _n_modules(7)
    assert "modules_aantal" not in _codes_in(check_rewrite(rwd, {**_CTX, "dagen": 5}), "modules")
    assert "modules_aantal" in _codes_in(check_rewrite(rwd, {**_CTX, "dagen": 3}), "modules", HARD)
    assert "modules_aantal" in _codes_in(check_rewrite(rwd, {**_CTX, "dagen": 1}), "modules", HARD)


def test_vijf_modules_is_schoon_op_elke_duur():
    rwd = _good_rewrite()
    rwd["modules"]["modules"] = _n_modules(5)
    for dagen in (1, 2, 3, 5):
        assert "modules_aantal" not in _codes_in(
            check_rewrite(rwd, {**_CTX, "dagen": dagen}), "modules")


def test_fewshot_toont_de_modulestructuur_met_niveaus():
    """Zonder nesting ziet de schrijver bij het zwaarst wegende kopje een platte lijst."""
    tekst = rw.goud_voorbeelden()
    modules_blok = tekst.split("**Modules**")[1].split("**Doelen**")[0]
    assert "\n* " in modules_blok, "geen moduletitels in het voorbeeld"
    assert "\n  * " in modules_blok, "geen sub-bullets in het voorbeeld"


def test_fewshot_noemt_de_duur_van_de_training_niet():
    import rewrite_checks as c
    tekst = rw.goud_voorbeelden(n=len(rw.GOUD_VOORBEELDEN))
    treffer = c._DUUR_RE.search(tekst)
    assert not treffer, f"few-shot noemt de duur: {treffer.group(0)!r}"


def test_fewshot_haalt_zelf_alle_harde_checks():
    """Een few-shot die de eigen regels schendt is erger dan geen few-shot.

    Dit is de enige test die het voorbeeldmateriaal zelf toetst. Verandert er een regel,
    dan valt hij hier om en niet pas in de output van de volgende batch.
    """
    bestanden = rw.goud_bestanden(rw.GOUD_V2_DIR)
    assert bestanden, "geen voorbeeldmateriaal in goud_v2"
    for pad in bestanden:
        with open(pad, encoding="utf-8") as f:
            d = json.load(f)
        titel = d.get("titel", "")
        concept = rw.goud_naar_check_input(d.get("content") or {}, titel)
        hard = hard_fails(check_rewrite(concept, {"naam": titel}))
        assert not hard, f"{os.path.basename(pad)}: {[str(i) for i in hard]}"


def test_fewshot_demonstreert_geen_verouderde_vaste_tekst():
    tekst = rw.goud_voorbeelden(n=len(rw.GOUD_VOORBEELDEN))
    assert tekst, "few-shot is leeg"
    for vervallen in sjabloon.VERVALLEN_VASTE_TEKSTEN:
        assert vervallen not in tekst, f"few-shot toont verouderde vaste tekst: {vervallen}"
    # en demonstreert wél het lerende aspect uit Sectie 0.15
    assert "kunnen" in tekst


def test_scan_vorm_ziet_verouderde_vaste_tekst():
    """Een training met de vorige generatie boilerplate mag nooit `overnemen` halen."""
    scan = rw.scan_vorm(_OUDE_CONTENT, "Training XML")
    assert scan["ondergrens"] == "format"
    assert scan["verouderde_vaste_tekst"]
    assert "vaste sjabloonteksten" in scan["reden"]


# ---------------------------------------------------------------------------
# Reviewronde 4: em-dash, reikwijdte, groep-intro's, nadruk in de Aanpak, goud promoveren
# ---------------------------------------------------------------------------

def test_em_dash_is_hard_in_elk_schrijversveld():
    for veld, waarde in (("overzicht", "Wil je leren — echt leren — hoe dit werkt?"),
                         ("inleiding", "Je werkt met data – en met modellen."),
                         ("doelgroep", "Deze training is bedoeld voor iedereen — echt iedereen.")):
        hard = checks.hard_fails(checks.check_em_dash({veld: waarde}))
        assert [i.code for i in hard] == ["em_dash"], f"{veld} laat de em-dash door"
    # het gewone koppelteken blijft gewoon toegestaan
    assert not checks.check_em_dash({"overzicht": "Je leert data-analyse hands-on toepassen."})


def test_geen_liggend_streepje_in_de_promptbestanden():
    """Wat de schrijver in zijn context ziet, schrijft hij ook.

    De em-dash is verboden én een harde check, maar de spec-bestanden stonden er zelf vol mee:
    173 stuks over vijf bestanden. Dat is precies het soort tegenstrijdigheid waar een model
    de verkeerde kant van kiest. Ze staan er nu uitgeschreven ("[liggend streepje]") in plaats
    van letterlijk, en deze test houdt dat zo.
    """
    for pad in (rw.SCHRIJFSPEC, rw.HUMANISERING, rw.STIJLREGISTER, rw.CORRECTIES,
                rw.BEOORDELINGSSPEC, rw.TEMPLATE_PATH):
        with open(pad, encoding="utf-8") as f:
            tekst = f.read()
        gevonden = [r for r in ("—", "–") if r in tekst]
        assert not gevonden, f"{os.path.basename(pad)} bevat {gevonden}"


def test_geen_liggend_streepje_in_de_systemprompts():
    """Ook de kop boven de few-shot en het goud zelf tellen mee.

    De bestandstest hierboven dekt de spec's; dit dekt wat de code eromheen bouwt, inclusief
    de voorbeelden. Een voorbeeld met een em-dash zou het teken tonen op precies de plek waar
    de schrijver naar vorm zoekt.
    """
    for naam, blokken in (("writer", rw.build_writer_system()),
                          ("judge", rw.build_judge_system()),
                          ("modus", rw.build_modus_system())):
        tekst = "".join(b["text"] for b in blokken)
        gevonden = [r for r in ("—", "–") if r in tekst]
        assert not gevonden, f"{naam}-systemprompt bevat {gevonden}"


def test_geen_liggend_streepje_in_de_userberichten_en_de_tools():
    """Het gat dat de twee tests hierboven lieten liggen.

    De systemprompt was schoon en de spec-bestanden ook, maar daaronder stond het teken er
    34 keer: 9 in de beschrijvingen van `submit_rewrite` (de tekst die de schrijver leest op het
    moment dat hij een kopje schrijft), 8 in het user-bericht van de schrijver, 8 in dat van de
    judge en de rest in de modus- en actualiseerprompts. De HARD-boodschap van `check_em_dash`
    deed het zelf ook: die gaat via `notes` letterlijk terug naar de schrijver, dus het teken
    stond in de zin die het verbood.

    De user-berichten wisselen per training, dus dit test wat de code eromheen bouwt met een
    minimale briefing. Wat uit de bron of uit een scorer-veld komt telt niet mee, dat kunnen we
    hier niet afdwingen; het scoringsproject houdt de kern schoon (`zonder_liggend_streepje`).
    """
    b = _briefing(huidige_content={"algemene_omschrijving": "Bestaande tekst."})
    stukken = {
        "writer-user": rw.build_writer_user(b),
        "judge-user": rw.build_judge_user(b, {"titel": "T", "overzicht": "Wil je iets?"}),
        "actualisatie-user": rw.build_actualisatie_user(b, b.huidige_content, "Training T"),
        "hergenereer-tool": json.dumps(rw.build_kopje_tool("overzicht"), ensure_ascii=False),
        "actualisatie-tool": json.dumps(rw.build_actualisatie_tool(), ensure_ascii=False),
        "modus-instructie": rw.SCHAT_MODUS_INSTRUCTIE,
        "kies-vervolg": rw.KIES_VERVOLG_SYSTEM,
    }
    for naam, tool in (("submit_rewrite", rw.SUBMIT_REWRITE), ("submit_judgment", rw.SUBMIT_JUDGMENT),
                       ("submit_modus", rw.SUBMIT_MODUS),
                       ("submit_vervolgstappen", rw.SUBMIT_VERVOLGSTAPPEN)):
        stukken[naam] = json.dumps(tool, ensure_ascii=False)
    for naam, tekst in stukken.items():
        gevonden = [r for r in ("—", "–") if r in tekst]
        assert not gevonden, f"{naam} bevat {gevonden}"

    # De correctie die de schrijver terugkrijgt mag het teken evenmin tonen.
    issues = checks.check_em_dash({"overzicht": "Wil je leren — echt leren?"})
    bericht = str(issues[0])
    assert "—" not in bericht and "–" not in bericht, f"de HARD-boodschap toont het teken: {bericht}"


def test_em_dash_check_raakt_de_vaste_teksten_niet():
    """De boilerplate komt nooit langs deze check; anders faalt elke training voor altijd."""
    rwin = {"overzicht": "Schoon.", "aanpak_invulling": "je dit toepast"}
    assert not checks.check_em_dash(rwin)
    for tekst in sjabloon.VASTE_TEKSTEN:
        assert "—" not in tekst and "–" not in tekst, f"vaste tekst bevat een em-dash: {tekst[:40]}"


def test_reikwijdte_flagt_een_voorwaarde_met_een_opsomming():
    """Uit 3127: de bron toont breedte, het concept maakte er een afbakening van."""
    smal = {"inleiding": "Werk je in communicatie, beleid, HR of klantcontact, dan ben je na "
                         "deze training in staat om AI in te zetten."}
    issues = checks.check_reikwijdte(smal)
    assert [i.code for i in issues] == ["reikwijdte"]
    assert not checks.hard_fails(issues), "dit is een leesbril, geen oordeel"
    # de insluitende vorm en een gewone voorwaardelijke zin blijven ongemoeid
    assert not checks.check_reikwijdte({"inleiding": "Of je nu in communicatie, beleid of HR "
                                                     "werkt, je leert prompts te schrijven."})
    assert not checks.check_reikwijdte({"voorkennis": "Werk je nog niet met SQL, dan volg je "
                                                      "eerst de basistraining."})


def test_eigen_case_inbrengen_is_hard():
    """Uit 3036: alleen een bedrijfstraining kan materiaal van de deelnemer verwerken."""
    for veld, waarde in (
            ("modules", "Een eigen veranderopgave rond datamanagement inbrengen"),
            ("inleiding", "De training biedt ruimte om eigen vraagstukken mee te nemen."),
            ("overzicht", "Je brengt je eigen case in en werkt die uit."),
            ("aanpak_invulling", "je je eigen dataset meeneemt naar de training"),
            ("doelen", "Je eigen praktijkopdracht aan te leveren en uit te werken"),
            ("inleiding", "Je legt je eigen probleemstelling voor aan de trainer.")):
        rwin = {veld: [waarde]} if veld == "doelen" else {veld: waarde}
        if veld == "modules":
            rwin = {"modules": {"modules": [{"titel": "M", "bullets": [waarde]}]}}
        hard = checks.hard_fails(checks.check_eigen_case(rwin))
        assert [i.code for i in hard] == ["eigen_case_inbrengen"], f"{veld}: {waarde}"


def test_bezittelijk_woord_op_een_case_is_hard():
    """"een praktijkcase" leveren wij; "je eigen praktijkcase" belooft er een van de deelnemer."""
    for waarde in ("Je past alles toe op je eigen praktijkcase.",
                   "Je sluit af met je eigen praktijkcase.",
                   "Je werkt met jullie eigen casussen.",
                   "Jouw praktijkcase"):
        hard = checks.hard_fails(checks.check_eigen_case({"overzicht": waarde}))
        assert [i.code for i in hard] == ["eigen_case"], waarde
    assert not checks.check_eigen_case(
        {"overzicht": "Je past alles toe op een praktijkcase."})


def test_eigen_case_ziet_door_koppeltekens_heen():
    """Een koppelteken is geen woordteken; zonder `[\\w-]` glipte de samenstelling erlangs.

    Kwam aan het licht doordat twee few-shot-voorbeelden (3127, 796) de fout zelf
    demonstreerden terwijl de check zweeg. Een gat in een check die de few-shot moet bewaken
    is duurder dan elders: het voorbeeld leert de fout aan.
    """
    for waarde in ("Richtlijnen opstellen en je eigen AI-toepassingscasus uitwerken",
                   "Je use-case uitwerken tot een toepasbare oplossing"):
        hard = checks.hard_fails(checks.check_eigen_case({"modules": {"modules": [
            {"titel": "M", "bullets": [waarde]}]}}))
        assert [i.code for i in hard] == ["eigen_case"], waarde


def test_eigen_case_dekt_document_proces_en_werkvraag():
    """Uit 3127: niet elk werkmateriaal heet een case."""
    waarde = "Je eigen document, proces of werkvraag inbrengen en uitwerken"
    hard = checks.hard_fails(checks.check_eigen_case({"modules": {"modules": [
        {"titel": "M", "bullets": [waarde]}]}}))
    assert [i.code for i in hard] == ["eigen_case_inbrengen"], waarde


def test_eigen_case_vangt_ook_de_herkomst():
    """Het bezit kan een zelfstandig naamwoord verderop staan: "casussen uit je eigen praktijk"."""
    for waarde in ("Je werkt aan casussen uit je eigen praktijk.",
                   "Je werkt aan casussen die je voorbereidt met materiaal uit je eigen werk."):
        hard = checks.hard_fails(checks.check_eigen_case({"inleiding": waarde}))
        assert [i.code for i in hard] == ["eigen_case_herkomst"], waarde
    # de herkomst zonder bezit is precies wat we wél leveren
    assert not checks.check_eigen_case(
        {"inleiding": "We gebruiken herkenbare casussen uit de praktijk."})


def test_eigen_case_laat_toe_wat_de_deelnemer_zelf_maakt():
    """Wij leveren de case; wat de deelnemer daarmee bouwt is de oefening en mag beloofd worden.

    Dit is de grens die het makkelijkst meeschuift zodra iemand de patronen verbreedt, en hij is
    bij de review expliciet scherpgesteld: bij 796 wordt niets ingebracht, want de praktijkcase
    komt van ons en de applicatie is het resultaat. De richting beslist, niet het woord "eigen".
    Zie schrijfspec Sectie 0.25 en Sectie 0.15.
    """
    for waarde in (
            # de letterlijke zin uit 796, het schoolvoorbeeld van de toegestane kant
            "Je leert de belangrijkste patronen te benoemen, zelfstandig toe te passen en in een "
            "praktijkcase te verwerken tot een eigen applicatie.",
            "Praktijkcase: een eigen applicatie ontwikkelen",
            "Een eigen applicatie ontwerpen waarin je design patterns toepast",
            "Een roadmap opstellen voor een SIEM-oplossing binnen je eigen organisatie",
            "Een passende monitoring-strategie voor je eigen AWS-omgeving op te stellen",
            "Hierdoor ben je in staat om OpenCV in te zetten in je eigen projecten."):
        assert not checks.check_eigen_case({"inleiding": waarde}), waarde


def test_eigen_case_laat_afstemming_op_de_werksituatie_staan():
    """De grens: aansluiten op jouw praktijk mag, materiaal meebrengen niet.

    Dat onderscheid is het hele punt van de check, en de eerste zin hieronder staat bijna
    letterlijk in `sjabloon.AANPAK_ALINEA_1`. Zou die vuren, dan faalt elke training.
    """
    for waarde in ("Er is veel ruimte voor jouw vragen en werksituatie.",
                   "Er is ruimte voor vragen en het inbrengen van eigen situaties.",
                   "Je vertaalt het geleerde naar je eigen organisatie.",
                   "Je leert eigen datavraagstukken gestructureerd te analyseren.",
                   "Je brengt de samenhang tussen kanalen en touchpoints in kaart."):
        assert not checks.check_eigen_case({"inleiding": waarde}), waarde
    for tekst in sjabloon.VASTE_TEKSTEN:
        assert not checks.check_eigen_case({"inleiding": tekst}), f"vaste tekst: {tekst[:50]}"


def test_groep_met_een_titel_valt_weg_en_de_rest_blijft():
    groepen = [{"intro": "Verdiepen:", "titels": ["A", "B", "C"]},
               {"intro": "Verbreden:", "titels": ["D"]}]
    assert rw.snoei_groepen(groepen) == [groepen[0]]
    # blijven er te weinig titels over, dan vervallen de groepen helemaal en valt de
    # weergave terug op één vlakke lijst
    assert rw.snoei_groepen([{"intro": "Verdiepen:", "titels": ["A", "B"]},
                             {"intro": "Verbreden:", "titels": ["C"]}]) == []


def test_groep_met_een_titel_haalt_geen_enkele_weergave():
    doc = _document()
    doc["vervolgstappen"]["groepen"] = [
        {"intro": "Verdiepen:", "titels": ["Training Power BI", "Training DAX"]},
        {"intro": "Verbreden:", "titels": ["Training T-SQL"]},
    ]
    html = uit.document_to_content(doc, {})["follow_up"]
    md = uit.render_markdown(doc, "Training Data")
    for weergave in (html, md):
        assert "Verdiepen:" in weergave
        assert "Verbreden:" not in weergave, "een intro met één training hoort nergens te staan"


def test_groep_met_een_titel_is_een_flag_geen_hardfail():
    rwin = {"vervolgstappen_groepen": [{"intro": "Verbreden:", "titels": ["A"]}]}
    issues = checks.check_vervolgstappen(rwin)
    assert [i.code for i in issues] == ["groep_te_klein"]
    assert not checks.hard_fails(issues), "de schrijver schrijft deze groepen niet"


def test_aanpak_nadruk_staat_in_beide_uitvoervormen():
    """Het template benadrukt twee delen; markdown én CMS-HTML horen dat te tonen.

    Sinds augustus 2026 met enkele aanhalingstekens in plaats van cursief: de site en de
    leerportalen gaven <em> niet goed weer. De nadruk zit nu in de tekst zelf, dus geen enkele
    weergave kan hem nog kwijtraken.
    """
    assert "‘kennis’" in sjabloon.AANPAK_ALINEA_2
    assert "‘toepassing binnen jouw organisatie en werksituatie’" in sjabloon.AANPAK_ALINEA_2
    assert "*" not in sjabloon.AANPAK_ALINEA_2

    doc = _document()
    html = uit.document_to_content(doc, {})["setup"]
    assert "<em>" not in html
    assert "‘kennis’" in html
    assert "‘toepassing binnen jouw organisatie en werksituatie’" in html
    # de eerste "kennis" ("de meest actuele kennis") blijft zonder aanhalingstekens
    assert "actuele kennis, maar" in html

    md = uit.render_markdown(doc, "Training Data").split("## Aanpak")[1]
    assert md.count("‘kennis’") == 1


def test_aanpak_van_voor_de_wissel_krijgt_alsnog_aanhalingstekens():
    """De 32 documenten op schijf dragen `*...*`; een herrender mag geen sterretjes tonen."""
    html = uit.render_aanpak("Een vertaalslag van *kennis* naar *toepassing*.")
    assert "‘kennis’" in html and "‘toepassing’" in html
    assert "*" not in html and "<em>" not in html


def test_aanpak_html_ontstaat_niet_uit_brontekst():
    """Escapen gaat vóór het omzetten van de markering; anders smokkelt de bron tags binnen."""
    html = uit.render_aanpak("Een <script>-tag en een *cursief* stuk.")
    assert "&lt;script&gt;" in html
    assert "‘cursief’" in html


def _goud_kandidaat(tid: str = "goed") -> dict:
    """Een artefact zoals `schrijf_training_artefacten` het wegschrijft, dat alles haalt."""
    writer_out = _good_rewrite()
    document = rw.assemble_document(writer_out, _briefing(), ["Training Power BI"])
    return {"training_id": tid, "titel": document["titel"], "status": rw.APPROVED,
            "document": document, "writer_out": writer_out,
            "content": uit.document_to_content(document, {"days": 2})}


def test_promoveer_naar_goud_kijkt_ook_in_de_batch_submappen():
    """De few-shot wordt gekozen uit álles wat we ooit schreven, niet uit de platte map alleen."""
    with tempfile.TemporaryDirectory() as tmp:
        bron, goud = os.path.join(tmp, "trainingen"), os.path.join(tmp, "goud")
        os.makedirs(os.path.join(bron, "ronde 3"))
        artefact = _goud_kandidaat("in_submap")
        with open(os.path.join(bron, "ronde 3", "in_submap.json"), "w", encoding="utf-8") as f:
            json.dump(artefact, f)
        uitslag = rw.promoveer_naar_goud(bron, goud, dry_run=True, verbose=False)
    assert [g["training_id"] for g in uitslag["gepromoveerd"]] == ["in_submap"]


def test_promoveer_naar_goud_kiest_alleen_wat_alle_checks_haalt():
    """Een training met een harde fail komt de goudmap niet in -- ook niet als hij approved is."""
    goed = _goud_kandidaat()
    fout = _goud_kandidaat("fout")
    fout["writer_out"]["overzicht"] = fout["document"]["overzicht"] = (
        "Wil je dit — echt dit — kunnen doen?")

    with tempfile.TemporaryDirectory() as tmp:
        bron, goud = os.path.join(tmp, "trainingen"), os.path.join(tmp, "goud")
        os.makedirs(bron)
        for artefact in (goed, fout):
            with open(os.path.join(bron, f"{artefact['training_id']}.json"), "w",
                      encoding="utf-8") as f:
                json.dump(artefact, f)
        uitslag = rw.promoveer_naar_goud(bron, goud, verbose=False)

        assert [g["training_id"] for g in uitslag["gepromoveerd"]] == ["goed"]
        assert [tid for tid, _ in uitslag["afgewezen"]] == ["fout"]
        assert os.path.exists(os.path.join(goud, "goed.json"))
        assert not os.path.exists(os.path.join(goud, "fout.json"))
        assert rw.laad_goud_selectie(goud) == ("goed",)


def test_promoveer_naar_goud_schrijft_niets_bij_dry_run():
    artefact = _goud_kandidaat()
    with tempfile.TemporaryDirectory() as tmp:
        bron, goud = os.path.join(tmp, "trainingen"), os.path.join(tmp, "goud")
        os.makedirs(bron)
        with open(os.path.join(bron, "goed.json"), "w", encoding="utf-8") as f:
            json.dump(artefact, f)
        uitslag = rw.promoveer_naar_goud(bron, goud, dry_run=True, verbose=False)
        assert [g["training_id"] for g in uitslag["gepromoveerd"]] == ["goed"]
        assert not os.path.exists(goud)


def test_goud_selectie_valt_terug_op_het_handwerk_zonder_manifest():
    """`herschreven/` staat in .gitignore: na een verse checkout is er geen manifest."""
    with tempfile.TemporaryDirectory() as tmp:
        assert rw.laad_goud_selectie(tmp) == rw._GOUD_V2_FALLBACK


# ---------------------------------------------------------------------------
# De wachtrij: welke trainingen draaien er, en waarom valt de rest af
# ---------------------------------------------------------------------------

def _wachtrij_situatie(d, ids, klaar=(), herschreven=None, mislukt=()):
    """Schrijft een scoresheet met `ids` in die volgorde, plus een herschreven.xlsx met `klaar`.

    `herschreven` is de kolom die zonder `modus_voorstel` de modus bepaalt: 1 -> `overnemen`.
    `mislukt` komt als `error`-rij in het review-blad: wél in het sheet, geen tekst erachter.
    """
    import pandas as pd
    scored_pad = os.path.join(d, "prio.xlsx")
    n = len(ids)
    _scored_df(id=list(ids), name=[f"Training {t}" for t in ids],
               herschreven=list(herschreven or [0] * n),
               actualiteit_actie=[""] * n, actie_besluit=[""] * n
               ).to_excel(scored_pad, index=False)
    if klaar or mislukt:
        os.makedirs(os.path.join(d, "uit"), exist_ok=True)
        rijen = ([{"training_id": t, "status": rw.APPROVED} for t in klaar]
                 + [{"training_id": t, "status": "error"} for t in mislukt])
        with pd.ExcelWriter(os.path.join(d, "uit", "herschreven.xlsx")) as w:
            pd.DataFrame(rijen).to_excel(w, sheet_name="review", index=False)
    return scored_pad, os.path.join(d, "uit")


def test_wachtrij_telt_na_de_skip_filters_en_niet_over_het_sheet():
    """De verwarring die deze functie bestaat om weg te nemen.

    `start` telde altijd al over de gefilterde lijst, maar dat was nergens te zien: met 10 van
    de 15 trainingen al klaar leverde START=3 de vierde van de vijf resterende op, niet
    sheetrij 3. De preview hoort beide nummeringen naast elkaar te zetten.
    """
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12, 13, 14], klaar=[10, 11])
        q = rw.bouw_wachtrij(scored, uit_dir, start=1, limit=1)
        gekozen = q[q["geselecteerd"]]
        assert list(gekozen["training_id"]) == [13], list(gekozen["training_id"])
        assert int(gekozen["sheet"].iloc[0]) == 3 and int(gekozen["wachtrij"].iloc[0]) == 1
        assert list(q[q["wachtrij"].notna()]["training_id"]) == [12, 13, 14]


def test_wachtrij_noemt_per_rij_waarom_hij_niet_meedraait():
    """Zonder reden per rij blijft "waarom draait deze niet" een raadsel dat je zelf natelt."""
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12], klaar=[10])
        q = rw.bouw_wachtrij(scored, uit_dir, start=0, limit=1).set_index("training_id")
        assert q.loc[10, "reden"] == "staat al in herschreven.xlsx"
        assert q.loc[11, "reden"] == ""          # draait
        assert q.loc[12, "reden"] == "buiten START/N"


def test_alleen_ids_negeert_start_en_limit():
    """Een id-selectie is het antwoord op een wachtrij die per run van lengte verandert."""
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12, 13])
        q = rw.bouw_wachtrij(scored, uit_dir, start=2, limit=1, alleen_ids=[10, 13])
        assert list(q[q["geselecteerd"]]["training_id"]) == [10, 13]


def test_alleen_ids_meldt_wat_er_niet_in_de_wachtrij_zat():
    """Stil minder draaien dan gevraagd is precies het gedrag dat hier wordt weggenomen."""
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11], klaar=[10])
        q = rw.bouw_wachtrij(scored, uit_dir, alleen_ids=[10, 99])
        assert q.attrs["ids_niet_gedraaid"] == [10, 99]


def test_limit_nul_selecteert_niets_en_niet_alles():
    """`if limit:` liet N=0 samenvallen met N=None -- de duurste betekenis van een nul."""
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12])
        assert not rw.bouw_wachtrij(scored, uit_dir, limit=0)["geselecteerd"].any()
        assert rw.bouw_wachtrij(scored, uit_dir, limit=None)["geselecteerd"].all()


def test_overnemen_rijen_staan_los_van_start_en_limit():
    """Het overnemen-spoor heeft zijn eigen lus in `rewrite_file`; start/limit raken het niet."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12], herschreven=[1, 0, 0])
        q = rw.bouw_wachtrij(scored, uit_dir, start=1, limit=1).set_index("training_id")
        assert q.loc[10, "spoor"] == "overnemen" and q.loc[10, "geselecteerd"]
        assert pd.isna(q.loc[10, "wachtrij"]), "een overnemen-rij hoort geen wachtrijpositie te krijgen"
        assert list(q[q["wachtrij"].notna()].index) == [11, 12]
        assert q.loc[12, "geselecteerd"] and not q.loc[11, "geselecteerd"]


def test_een_error_rij_blokkeert_de_volgende_run_niet():
    """Anders is "één fout kost de batch" verruild voor "één fout kost stil die training".

    De rij staat in `herschreven.xlsx`, maar er ligt geen tekst achter: hervatten hoort hem
    dus opnieuw aan te dragen. Een geslaagde rij blijft wél overgeslagen.
    """
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12], klaar=[10], mislukt=[11])
        q = rw.bouw_wachtrij(scored, uit_dir).set_index("training_id")
        assert q.loc[10, "reden"] == "staat al in herschreven.xlsx"
        assert q.loc[11, "geselecteerd"] and q.loc[12, "geselecteerd"]


def test_review_blad_zonder_statuskolom_telt_nog_steeds_als_klaar():
    """Sheets van vóór de error-route hebben geen `status`; die mogen niet ineens herdraaien."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11])
        os.makedirs(uit_dir, exist_ok=True)
        with pd.ExcelWriter(os.path.join(uit_dir, "herschreven.xlsx")) as w:
            pd.DataFrame({"training_id": [10]}).to_excel(w, sheet_name="review", index=False)
        q = rw.bouw_wachtrij(scored, uit_dir).set_index("training_id")
        assert q.loc[10, "reden"] == "staat al in herschreven.xlsx"
        assert q.loc[11, "geselecteerd"]


def test_een_stukgelopen_training_kost_niet_de_hele_batch():
    """De duurste fout van batch 1: training 1 gooide, en alle 46 waren weg.

    `herschreven.xlsx` wordt pas ná de lus geschreven, dus een uitzondering halverwege laat
    niets achter -- ook niet voor de trainingen die het wél haalden. Eén fout hoort één
    `error`-rij te kosten, met de uitzondering in de kolom `reden`.
    """
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12])
        bron = os.path.join(d, "bron.xlsx")
        pd.DataFrame([{"id": t, "name": f"Training {t}",
                       "content": json.dumps(_content(), ensure_ascii=False)}
                      for t in (10, 11, 12)]).to_excel(bron, index=False)
        besluiten = os.path.join(d, "besluiten.xlsx")   # geen acties, dus geen rijen
        pd.DataFrame(columns=["training_id", "nr", "actie", "besluit"]).to_excel(
            besluiten, index=False)

        def _valt_om_bij_11(client, b, catalog, boom=None):
            if b.training_id == 11:
                raise ValueError("Streaming is required for operations")
            return rw.RewriteResult(b.training_id, b.nieuwe_titel, rw.APPROVED, modus=b.modus)

        echte_client, echte_rewrite = rw.make_client, rw.rewrite_one
        rw.make_client, rw.rewrite_one = (lambda: None), _valt_om_bij_11
        try:
            review = rw.rewrite_file(scored, bron, uit_dir, besluiten_path=besluiten,
                                     verbose=False)
        finally:
            rw.make_client, rw.rewrite_one = echte_client, echte_rewrite

        rijen = review.set_index("training_id")
        assert list(rijen["status"]) == [rw.APPROVED, "error", rw.APPROVED], list(rijen["status"])
        assert "ValueError" in rijen.loc[11, "reden"]
        # en het sheet staat op schijf, inclusief de twee die het wél haalden
        opnieuw = pd.read_excel(os.path.join(uit_dir, "herschreven.xlsx"), sheet_name="review")
        assert sorted(opnieuw["training_id"]) == [10, 11, 12]


def test_een_verstreken_tijdsbudget_kost_die_ene_training_en_niet_de_batch():
    """Waar het budget voor bestaat: de vastgelopen training eruit, de rest gewoon door.

    `TijdOverschreden` moet daarvoor een gewone `Exception` zijn. Erft hij van BaseException,
    dan vangt de lus hem niet en kost één trage training alsnog alle 46 -- ook de trainingen
    die al klaar waren, want `herschreven.xlsx` wordt pas ná de lus geschreven. En omdat de
    rij een `error` is, plant `bouw_wachtrij` hem bij de volgende run vanzelf opnieuw in.
    """
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11, 12])
        bron = os.path.join(d, "bron.xlsx")
        pd.DataFrame([{"id": t, "name": f"Training {t}",
                       "content": json.dumps(_content(), ensure_ascii=False)}
                      for t in (10, 11, 12)]).to_excel(bron, index=False)
        besluiten = os.path.join(d, "besluiten.xlsx")
        pd.DataFrame(columns=["training_id", "nr", "actie", "besluit"]).to_excel(
            besluiten, index=False)

        def _loopt_vast_bij_11(client, b, catalog, boom=None):
            if b.training_id == 11:
                raise rw.TijdOverschreden("tijdsbudget van 25 minuten verstreken")
            return rw.RewriteResult(b.training_id, b.nieuwe_titel, rw.APPROVED, modus=b.modus)

        echte_client, echte_rewrite = rw.make_client, rw.rewrite_one
        rw.make_client, rw.rewrite_one = (lambda: None), _loopt_vast_bij_11
        try:
            review = rw.rewrite_file(scored, bron, uit_dir, besluiten_path=besluiten,
                                     verbose=False)
        finally:
            rw.make_client, rw.rewrite_one = echte_client, echte_rewrite

        rijen = review.set_index("training_id")
        assert list(rijen["status"]) == [rw.APPROVED, "error", rw.APPROVED], list(rijen["status"])
        assert "tijdsbudget" in rijen.loc[11, "reden"]
        # de duur staat erbij, ook (juist) van de training die het niet haalde
        assert set(review["seconden"] >= 0) == {True}
        # en 11 draait de volgende run gewoon weer mee
        q = rw.bouw_wachtrij(scored, uit_dir).set_index("training_id")
        assert q.loc[11, "geselecteerd"] and not q.loc[10, "geselecteerd"]


def test_start_voorbij_het_einde_waarschuwt_in_plaats_van_stil_niets_te_doen():
    """Een run die 0 trainingen draait zonder één regel uitvoer leest als een geslaagde run."""
    with tempfile.TemporaryDirectory() as d:
        scored, uit_dir = _wachtrij_situatie(d, [10, 11], klaar=[10])
        q = rw.bouw_wachtrij(scored, uit_dir, start=5)
        assert not q["geselecteerd"].any()
        melding = " ".join(rw._wachtrij_waarschuwingen(q, 5, None, None))
        assert "START=5" in melding and "1 trainingen" in melding, melding


# ---------------------------------------------------------------------------
# Upload naar Google Drive
# ---------------------------------------------------------------------------

class _DriveFout(Exception):
    """Staat voor een HttpError; die klasse importeren zou googleapiclient vereisen."""


def _upload_html(kw) -> str:
    """De HTML die als media aan `files.create`/`files.update` is meegegeven."""
    return kw["media_body"].getbytes(0, kw["media_body"].size()).decode("utf-8")


def _fake_drive(bestaand=(), faal_op=(), comments_stuk=False, vreemde_comments=()):
    """Een Drive-service zonder netwerk, in hetzelfde SimpleNamespace-idioom als de fake client.

    `bestaand` zijn de files die `files().list` teruggeeft (op naam gefilterd), `faal_op` de
    namen waarop `create` een fout gooit -- nodig om te laten zien dat één mislukte upload de
    rest van de batch niet meeneemt. `comments_stuk` laat het plaatsen van de opmerking falen,
    wat een geslaagde upload niet mag omkatten naar een mislukte. `vreemde_comments` zijn
    opmerkingen van een reviewer die er al staan; die mag het opruimen nooit aanraken.
    """
    gemaakt, vervangen = [], []
    comments = [{"id": f"r{i}", "fileId": f, "content": t, "author": {"me": eigen}}
                for i, (f, t, eigen) in enumerate(vreemde_comments)]

    def _list(**kw):
        naam = kw.get("q", "")
        treffers = [f for f in bestaand if f"name = '{f['name']}'" in naam]
        return SimpleNamespace(execute=lambda **_: {"files": treffers})

    def _create(**kw):
        naam = kw["body"]["name"]
        if naam in faal_op:
            raise _DriveFout(f"503 op {naam}")
        gemaakt.append(kw)
        fid = f"f{len(gemaakt)}"
        return SimpleNamespace(execute=lambda **_: {
            "id": fid, "name": naam, "webViewLink": f"https://docs.google.com/d/{fid}"})

    def _update(**kw):
        vervangen.append(kw)
        return SimpleNamespace(execute=lambda **_: {
            "id": kw["fileId"], "name": "", "webViewLink": f"https://docs.google.com/d/{kw['fileId']}"})

    def _comment_create(**kw):
        if comments_stuk:
            raise _DriveFout("403 op de opmerking")
        comments.append({"id": f"c{len(comments)}", "fileId": kw["fileId"],
                         "content": kw["body"]["content"], "author": {"me": True}})
        return SimpleNamespace(execute=lambda **_: {"id": comments[-1]["id"]})

    def _comment_list(**kw):
        staand = [c for c in comments if c["fileId"] == kw["fileId"]]
        return SimpleNamespace(execute=lambda **_: {"comments": staand})

    def _comment_delete(**kw):
        weg = [c for c in comments if c["id"] == kw["commentId"]]
        for c in weg:
            comments.remove(c)
        return SimpleNamespace(execute=lambda **_: None)

    files = SimpleNamespace(list=_list, create=_create, update=_update)
    return SimpleNamespace(
        files=lambda: files,
        comments=lambda: SimpleNamespace(create=_comment_create, list=_comment_list,
                                         delete=_comment_delete),
        gemaakt=gemaakt, vervangen=vervangen, comments_op_docs=comments)


def _artefact(d, tid, titel="Training Data", batch=None, **overrides):
    """Eén <id>.json op schijf, zoals `schrijf_training_artefacten` hem wegschrijft."""
    map_ = rw.artefact_dir(d, batch)
    os.makedirs(map_, exist_ok=True)
    data = {"training_id": tid, "titel": titel, "status": rw.APPROVED,
            "content": _content(), **overrides}
    with open(os.path.join(map_, f"{tid}.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    return data


def test_docnaam_volgt_het_afgesproken_format():
    assert drive.docnaam(2347, "Training Big Data") == \
        "2347 - Training Big Data (automatisch herschreven)"


def test_docnaam_klapt_witruimte_in_de_titel_in():
    """Een regeleinde in de titel maakt een Drive-naam die je in de lijst niet meer leest."""
    assert drive.docnaam(27, "Training\nSQL  Basis") == \
        "27 - Training SQL Basis (automatisch herschreven)"


def test_upload_zet_de_conversie_mimetype_op_google_docs():
    """Doelformaat in de body, bronformaat in de media: dat verschil is de hele conversie."""
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
    docs = [c for c in service.gemaakt if c["body"]["mimeType"] == drive.DOC_MIME]
    assert len(docs) == 1, service.gemaakt
    assert docs[0]["body"]["name"] == "2347 - Training Data (automatisch herschreven)"
    assert docs[0]["media_body"].mimetype() == "text/html"
    assert docs[0]["fields"] == "id,name,webViewLink"


def test_upload_maakt_de_batchmap_onder_de_rootmap():
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        res = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                      verbose=False)
    mappen = [c for c in service.gemaakt if c["body"]["mimeType"] == drive.MAP_MIME]
    assert len(mappen) == 1, mappen
    assert mappen[0]["body"]["name"] == "batch 1" and mappen[0]["body"]["parents"] == ["root"]
    # de docs komen in de batchmap terecht, niet los in de rootmap
    docs = [c for c in service.gemaakt if c["body"]["mimeType"] == drive.DOC_MIME]
    assert docs[0]["body"]["parents"] == [res["map_id"]] != ["root"]


def test_batchmap_wordt_hergebruikt_in_plaats_van_opnieuw_aangemaakt():
    """Drive staat twee mappen met dezelfde naam toe; zonder zoektocht groeit de Drive vol."""
    service = _fake_drive(bestaand=[{"id": "m1", "name": "batch 1",
                                     "webViewLink": "https://drive/m1"}])
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        res = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                      verbose=False)
    assert res["map_id"] == "m1"
    assert not [c for c in service.gemaakt if c["body"]["mimeType"] == drive.MAP_MIME]


def test_upload_slaat_een_training_over_die_al_in_het_manifest_staat():
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        eerst = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                        verbose=False)
        opnieuw = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                          verbose=False)
    assert eerst["nieuw"] == [2347] and opnieuw["nieuw"] == []
    assert opnieuw["overgeslagen"] == [2347]
    assert opnieuw["urls"][2347] == eerst["urls"][2347]
    assert len([c for c in service.gemaakt if c["body"]["mimeType"] == drive.DOC_MIME]) == 1


def test_upload_vindt_een_bestaand_doc_op_naam_als_het_manifest_weg_is():
    """Anders levert een weggegooid manifest een tweede doc naast het eerste op."""
    naam = drive.docnaam(2347, "Training Data")
    service = _fake_drive(bestaand=[{"id": "d9", "name": naam, "webViewLink": "https://docs/d9"}])
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        res = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                      verbose=False)
    assert res["overgeslagen"] == [2347] and res["urls"][2347] == "https://docs/d9"
    assert not [c for c in service.gemaakt if c["body"]["mimeType"] == drive.DOC_MIME]


def test_upload_meldt_gewijzigde_tekst_in_plaats_van_stil_over_te_slaan():
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
        _artefact(d, 2347, content=_content(summary="Een heel andere samenvatting."))
        uitvoer = io.StringIO()
        with contextlib.redirect_stdout(uitvoer):
            drive.upload_naar_drive(d, "batch 1", service=service, root_id="root")
    assert "de tekst is gewijzigd" in uitvoer.getvalue(), uitvoer.getvalue()


def test_nieuwe_versie_werkt_het_bestaande_doc_bij():
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
        res = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                      bij_bestaand="nieuwe_versie", verbose=False)
    assert res["nieuw"] == [2347]
    assert [c["fileId"] for c in service.vervangen] == ["f2"]


def test_een_mislukte_upload_stopt_de_rest_van_de_batch_niet():
    faal = drive.docnaam(2, "Training Data")
    service = _fake_drive(faal_op=[faal])
    with tempfile.TemporaryDirectory() as d:
        for tid in (1, 2, 3):
            _artefact(d, tid)
        uitvoer = io.StringIO()
        with contextlib.redirect_stdout(uitvoer):
            res = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                          verbose=False)
        # het manifest houdt de geslaagde vast, zodat een herdraai alleen de rest oppakt
        manifest = drive.lees_manifest(d)
    assert res["nieuw"] == [1, 3] and res["mislukt"] == [2]
    assert sorted(manifest["docs"]) == ["1", "3"]
    assert "upload mislukt" in uitvoer.getvalue()


def test_verzamel_uit_map_slaat_een_artefact_zonder_content_over():
    """Een leeg doc ziet een reviewer aan voor werk dat af is."""
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 1)
        _artefact(d, 2, status="error", content={})
        gevonden = drive.verzamel_uit_map(d)
    assert [t["training_id"] for t in gevonden] == [1]


def test_verzamel_uit_map_respecteert_alleen_ids():
    with tempfile.TemporaryDirectory() as d:
        for tid in (1, 2, 3):
            _artefact(d, tid)
        assert [t["training_id"] for t in drive.verzamel_uit_map(d, [3, 1])] == [1, 3]


def test_een_tweede_batch_neemt_de_eerste_niet_mee():
    """Dit was de aanleiding voor de submappen: de map van batch 2 kreeg ook alles uit batch 1."""
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 1, batch="ronde 1")
        _artefact(d, 2, batch="ronde 1")
        eerst = drive.upload_naar_drive(d, "ronde 1", service=service, root_id="root",
                                        verbose=False)
        _artefact(d, 3, batch="ronde 2")
        tweede = drive.upload_naar_drive(d, "ronde 2", service=service, root_id="root",
                                         verbose=False)
    assert eerst["nieuw"] == [1, 2]
    assert tweede["nieuw"] == [3], "batch 2 sleepte de trainingen van batch 1 mee"


def test_zonder_batch_pakt_de_upload_de_submap_met_dezelfde_naam():
    """`upload_naar_drive(OUT_DIR, "ronde 3")` zonder batch mag niet stil de platte map pakken."""
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 99)                      # oude training, platte map
        _artefact(d, 1, batch="ronde 3")
        res = drive.upload_naar_drive(d, "ronde 3", service=service, root_id="root",
                                      verbose=False)
    assert res["nieuw"] == [1], "de platte map werd meegenomen in plaats van de submap"


def test_zonder_submap_blijft_de_platte_map_de_bron():
    """De trainingen van voor de indeling hoeven niet te verhuizen om geüpload te worden."""
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 99)
        res = drive.upload_naar_drive(d, "losse map", service=service, root_id="root",
                                      verbose=False)
    assert res["nieuw"] == [99]


def test_verzamel_uit_map_kijkt_niet_in_de_submappen():
    """Recursief zoeken zou elke Drive-map opnieuw alles geven; dat is wat we oplossen."""
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 99)
        _artefact(d, 1, batch="ronde 1")
        assert [t["training_id"] for t in drive.verzamel_uit_map(d)] == [99]
        assert [t["training_id"] for t in drive.verzamel_uit_map(d, batch="ronde 1")] == [1]


def test_verzamel_uit_map_sorteert_numeriek():
    """Op bestandsnaam zou 27 tussen 2669 en 2725 belanden."""
    with tempfile.TemporaryDirectory() as d:
        for tid in (2725, 27, 2669):
            _artefact(d, tid)
        assert [t["training_id"] for t in drive.verzamel_uit_map(d)] == [27, 2669, 2725]


def test_upload_weigert_een_lege_mapnaam():
    """Zonder mapnaam belandt de hele batch los in de rootmap; dat is niet terug te draaien."""
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 1)
        for leeg in ("", "   "):
            try:
                drive.upload_naar_drive(d, leeg, service=_fake_drive(), root_id="root")
                assert False, f"lege mapnaam {leeg!r} werd geaccepteerd"
            except ValueError:
                pass


def _herschreven_xlsx(pad, ids, extra=None):
    """Een `herschreven.xlsx` met alleen de kolommen die `_zet_drive_urls_in_xlsx` aanraakt."""
    import pandas as pd
    review = pd.DataFrame([{"training_id": t, "titel": f"T{t}", "brontekst": "x",
                            **(extra or {})} for t in ids])
    with pd.ExcelWriter(pad) as writer:
        pd.DataFrame([{"id": t, "name": f"T{t}", "content": "{}"} for t in ids]).to_excel(
            writer, sheet_name="cms", index=False)
        review.to_excel(writer, sheet_name="review", index=False)


def test_drive_url_komt_als_laatste_kolom_in_het_reviewblad():
    """Het plakblok van het gedeelde sheet ligt vast; wat wij erbij verzinnen komt erachter."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "herschreven.xlsx")
        _herschreven_xlsx(pad, [1, 2])
        rw._zet_drive_urls_in_xlsx(pad, {1: "https://docs/a"}, verbose=False)
        review = pd.read_excel(pad, sheet_name="review")
    assert list(review.columns)[-1] == "drive_url"
    assert list(review["drive_url"].fillna("")) == ["https://docs/a", ""]


def test_een_deelupload_wist_geen_bestaande_drive_links():
    """Batch 2 uploadt niet wat in batch 1 al stond; die links moeten blijven staan."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "herschreven.xlsx")
        _herschreven_xlsx(pad, [1, 2], extra={"drive_url": ""})
        rw._zet_drive_urls_in_xlsx(pad, {1: "https://docs/a"}, verbose=False)
        rw._zet_drive_urls_in_xlsx(pad, {2: "https://docs/b"}, verbose=False)
        review = pd.read_excel(pad, sheet_name="review")
    assert list(review["drive_url"]) == ["https://docs/a", "https://docs/b"]


def test_een_lege_drive_url_blijft_leeg_en_wordt_geen_nan():
    """Een lege cel komt uit Excel terug als NaN, en NaN is truthy -- vandaar `_cel`.

    Uitgelezen met openpyxl en niet met pandas: `read_excel` maakt van de tekst "nan" opnieuw
    een NaN, en dan is de fout precies zo onzichtbaar als in het sheet zelf niet.
    """
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "herschreven.xlsx")
        _herschreven_xlsx(pad, [1, 2], extra={"drive_url": ""})
        rw._zet_drive_urls_in_xlsx(pad, {1: "https://docs/a"}, verbose=False)
        blad = openpyxl.load_workbook(pad)["review"]
        cellen = [str(c.value) for c in blad["A"][1:]]   # training_id, zonder de kopregel
        kolom = [k for k, c in enumerate(blad[1], start=1) if c.value == "drive_url"][0]
        waarden = [blad.cell(row=r + 2, column=kolom).value for r in range(len(cellen))]
    assert "nan" not in [str(w) for w in waarden], waarden
    assert waarden[0] == "https://docs/a" and not (waarden[1] or "")


def test_drive_urls_laten_het_cms_tabblad_ongemoeid():
    """Het cms-tabblad gaat rechtstreeks het CMS in; daar hoort geen reviewkolom in."""
    import pandas as pd
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "herschreven.xlsx")
        _herschreven_xlsx(pad, [1])
        rw._zet_drive_urls_in_xlsx(pad, {1: "https://docs/a"}, verbose=False)
        bladen = pd.read_excel(pad, sheet_name=None)
    assert sorted(bladen) == ["cms", "review"]
    assert list(bladen["cms"].columns) == ["id", "name", "content"]


def test_wrapper_zet_de_links_in_het_sheet_na_de_upload():
    """De uploadmodule kent geen pandas; deze schil is de enige plek waar het sheet bijkomt."""
    import pandas as pd
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 1)
        _herschreven_xlsx(os.path.join(d, "herschreven.xlsx"), [1])
        rw.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
        review = pd.read_excel(os.path.join(d, "herschreven.xlsx"), sheet_name="review")
    assert review["drive_url"][0].startswith("https://docs.google.com/d/")


def test_zonder_drive_map_raakt_de_batch_drive_niet():
    """Geen mapnaam = geen OAuth, geen manifest, geen enkele call. De default blijft offline."""
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 1)
        rw._upload_na_batch(d, None, None, False)
        rw._upload_na_batch(d, "", None, False)
        assert not os.path.exists(drive.manifest_pad(d))


def test_een_kapotte_upload_sloopt_een_geslaagde_batch_niet():
    """De artefacten staan al op schijf; de melding moet naar de losse aanroep wijzen."""
    class _StukService:
        def files(self):
            raise RuntimeError("token verlopen")

    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 1)
        uitvoer = io.StringIO()
        with contextlib.redirect_stdout(uitvoer):
            rw._upload_na_batch(d, "batch 1", _StukService(), False)
    assert "uploaden naar Drive mislukt" in uitvoer.getvalue()
    assert "rw.upload_naar_drive(" in uitvoer.getvalue(), uitvoer.getvalue()


def test_reden_valt_terug_op_de_revisienotities_van_de_judge():
    """De reden dat iets naar een mens gaat is het interessantste veld, en juist dat was leeg.

    `human_reden` vult de judge alleen als hij zelf naar de mens routeert; blijft hij tot het
    eind bij `needs-revision`, dan bleef er "judge: needs-revision na max revisies" over -- 37
    tekens over de trainingen waar tweemaal herschrijven niet hielp. 5 van de 8 human-queue-rijen
    in batch 1 zagen er zo uit, terwijl `revisie_notities` het concrete oordeel bevatte.
    """
    judgment = {"verdict": "needs-revision", "human_reden": "",
                "revisie_notities": ["Modules: module 4 en 5 overlappen.",
                                     "Inleiding: 'plaatsen' staat aan de onderkant."]}
    reden = rw._reden_uit_revisies(judgment)
    assert "module 4 en 5 overlappen" in reden
    assert "- Inleiding:" in reden
    # human_reden wint als de judge hem wél invult
    assert rw._reden_uit_revisies({"revisie_notities": []}).startswith("judge: needs-revision")


def test_comment_toont_alleen_de_flags_die_om_een_oordeel_vragen():
    """Alles tonen zou hier hetzelfde doen als de oude verzamelkolom: dan leest niemand het."""
    training = {"status": rw.APPROVED, "flags": ["laag: te lang", "hoog: verzonnen feit"],
                "flags_tier": {checks.TIER_HOOG: ["hoog: verzonnen feit"],
                               checks.TIER_LAAG: ["laag: te lang"]}}
    tekst = drive.comment_tekst(training)
    assert "hoog: verzonnen feit" in tekst
    assert "laag: te lang" not in tekst
    assert "1 punt om op te letten" in tekst, tekst


def test_comment_valt_terug_op_alle_flags_zonder_tier():
    """Artefacten van vóór `flags_tier`: liever te veel tonen dan iets verstoppen."""
    tekst = drive.comment_tekst({"status": rw.APPROVED, "flags": ["a", "b"], "flags_tier": {}})
    assert "- a" in tekst and "- b" in tekst
    assert "2 punten om op te letten" in tekst


def test_comment_draagt_de_reden_voor_de_human_queue():
    """Het interessantste veld: de flags zeggen wat de code zag, de reden wat de judge zag."""
    tekst = drive.comment_tekst({"status": rw.HUMAN_QUEUE, "reden": "judge wees af",
                                 "flags": []})
    assert rw.HUMAN_QUEUE in tekst and "judge wees af" in tekst
    assert "geen opmerkingen" in tekst.lower()


def test_comment_zet_de_letterlijke_backslash_n_van_de_judge_om():
    """De judge levert zijn toelichting soms met de tékens backslash en n, 4x in training 7."""
    tekst = drive.comment_tekst({"status": rw.HUMAN_QUEUE, "flags": [],
                                 "reden": "Eerste deel.\\n\\nTweede deel."})
    assert "\\n" not in tekst
    assert "Eerste deel." in tekst and "Tweede deel." in tekst


def test_comment_zonder_flags_zegt_dat_ook():
    tekst = drive.comment_tekst({"status": rw.APPROVED, "flags": [], "flags_tier": {}})
    assert "geen opmerkingen" in tekst.lower()
    assert "\n- " not in tekst


def test_elk_nieuw_doc_krijgt_een_opmerking_met_de_flags():
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347, flags=["[FLAG] verzonnen feit"],
                  flags_tier={checks.TIER_HOOG: ["[FLAG] verzonnen feit"]})
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
    assert len(service.comments_op_docs) == 1, service.comments_op_docs
    gezet = service.comments_op_docs[0]
    assert gezet["fileId"] == "f2", gezet            # f1 is de batchmap
    assert "[FLAG] verzonnen feit" in gezet["content"]


def test_de_flags_staan_nooit_in_de_tekst_van_het_doc():
    """Ze horen in de opmerking; in het document zouden ze in het CMS-artefact meeliften."""
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347, flags=["[FLAG] verzonnen feit"],
                  flags_tier={checks.TIER_HOOG: ["[FLAG] verzonnen feit"]})
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
    html = _upload_html(service.gemaakt[-1])
    assert "[FLAG]" not in html and "Automatisch herschreven" not in html


def test_een_vervangen_doc_krijgt_een_verse_opmerking_en_geen_tweede():
    """Na `files.update` is de oude opmerking losgeslagen, en hij beschrijft de vorige versie.

    Docs kan een ankerloze opmerking na een inhoudswissel nergens meer plaatsen en toont hem in
    de geschiedenis onder "oorspronkelijke content verwijderd". Laten staan levert dus geen
    opmerking op maar een spoor; er moet er precies één zijn, en dat is de nieuwe.
    """
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347, flags=["oud punt"], flags_tier={checks.TIER_HOOG: ["oud punt"]})
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
        _artefact(d, 2347, flags=["nieuw punt"], flags_tier={checks.TIER_HOOG: ["nieuw punt"]})
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                bij_bestaand="nieuwe_versie", verbose=False)
    assert len(service.comments_op_docs) == 1, service.comments_op_docs
    assert "nieuw punt" in service.comments_op_docs[0]["content"]


def test_het_opruimen_raakt_de_opmerkingen_van_een_reviewer_niet():
    """Alleen wat wij zelf schreven mag weg; de reviewer is de reden dat het doc bestaat."""
    service = _fake_drive(vreemde_comments=[("f2", "Deze module klopt niet", False),
                                            ("f2", "Automatisch herschreven. oud", True)])
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root", verbose=False)
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                bij_bestaand="nieuwe_versie", verbose=False)
    inhoud = [c["content"] for c in service.comments_op_docs]
    assert "Deze module klopt niet" in inhoud, inhoud
    assert "Automatisch herschreven. oud" not in inhoud, inhoud


def test_een_mislukte_opmerking_maakt_de_upload_niet_ongeldig():
    """Het document staat er en is bruikbaar; alleen de opmerking ontbreekt."""
    service = _fake_drive(comments_stuk=True)
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        uitvoer = io.StringIO()
        with contextlib.redirect_stdout(uitvoer):
            res = drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                          verbose=False)
    assert res["nieuw"] == [2347] and res["mislukt"] == []
    assert "opmerking plaatsen mislukt" in uitvoer.getvalue()


def test_met_comment_uit_raakt_de_opmerkingen_niet():
    service = _fake_drive()
    with tempfile.TemporaryDirectory() as d:
        _artefact(d, 2347)
        drive.upload_naar_drive(d, "batch 1", service=service, root_id="root",
                                met_comment=False, verbose=False)
    assert service.comments_op_docs == []


def test_manifest_wordt_atomisch_geschreven():
    """Halveert het schrijven, dan uploadt de volgende run alles opnieuw."""
    with tempfile.TemporaryDirectory() as d:
        drive.schrijf_manifest(d, {"root_id": "r", "mappen": {}, "docs": {}})
        assert not [f for f in os.listdir(d) if f.endswith(".tmp")]
        assert drive.lees_manifest(d)["root_id"] == "r"


# ---------------------------------------------------------------------------
# Runner (zonder pytest)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    passed, failed = 0, 0
    for t in tests:
        try:
            t()
            passed += 1
            print(f"PASS  {t.__name__}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL  {t.__name__}: {e}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"ERROR {t.__name__}: {type(e).__name__}: {e}")
    print(f"\n{passed} passed, {failed} failed")
    raise SystemExit(1 if failed else 0)
